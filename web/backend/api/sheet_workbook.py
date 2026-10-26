"""The sheet editor's workbook model: .xlsx / .csv <-> Univer workbook JSON.

The dashboard's sheet editor is Univer (open-source core, Apache-2.0,
web/static/vendor/univer). Univer's own .xlsx import/export is in its paid
tier, so the conversion lives here, on openpyxl (MIT), and it is
deliberately one-sided about what it owns:

* **Read** (``read_workbook``) turns a file into Univer's ``IWorkbookData``:
  every sheet (name, order, hidden, tab colour), values, formulas, number
  formats, fonts, fills, alignment, borders, merges, column widths, row
  heights and frozen panes, inside a ``WINDOW_ROWS`` x ``WINDOW_COLS``
  window per sheet.
* **Write** (``write_workbook``) applies an edited snapshot back INTO the
  workbook that is on disk. A cell, row, column, merge or sheet property
  is only touched when the snapshot differs from what ``read_workbook``
  would have produced for it, so everything the editor does not model
  (comments, data validation, conditional formatting, print setup, theme
  colours, defined names, anything past the window) stays exactly as it
  was. A brand-new name starts from a blank workbook.

Sheet identity: a sheet read from the file is given the id ``s<index>``.
The snapshot names, per sheet, the ``origin`` id it came from (``None``
for a sheet added in the editor); after a save the file's sheets are in
the snapshot's order, so the next save's origins are ``s0..sN`` again.

.csv is one sheet of values: formatting is not saved (the editor hides
those controls for .csv), and the existing csv writer keeps everything
past its window.
"""

from __future__ import annotations

import csv
import datetime as _dt
import math
import re
from copy import copy
from pathlib import Path
from typing import Any, Optional

from fastapi import HTTPException

WINDOW_ROWS = 5000
WINDOW_COLS = 100

# What one workbook may be, checked before openpyxl is asked to do anything
# (security review round 1, 2026-10-26). openpyxl holds a cell in ~540
# bytes while saving, and a merge makes an object for every cell it
# covers, so these bound the one-thread "sheets" pool's time and memory:
# about 250k cells is ~135 MB and a few seconds.
MAX_SHEETS = 50
MAX_SAVE_CELLS = 250_000      # cells a save may carry, all sheets together
MAX_OPEN_CELLS = 200_000      # cells an editable open may send; past it: read-only
MAX_MERGES = 5_000            # merged ranges per sheet (~1 ms of openpyxl each)
MAX_MERGED_AREA = 100_000     # cells covered by merges, per sheet
MAX_STYLES = 20_000
MAX_TEXT = 32_767             # Excel's own limit for one cell
MAX_COL_PX = 5_000
MAX_ROW_PX = 2_000

# Univer enums (core/lib/types/types/enum): BorderStyleTypes, HorizontalAlign,
# VerticalAlign, WrapStrategy, CellValueType.
_BORDER_TO_UNIVER = {
    "thin": 1, "hair": 2, "dotted": 3, "dashed": 4, "dashDot": 5,
    "dashDotDot": 6, "double": 7, "medium": 8, "mediumDashed": 9,
    "mediumDashDot": 10, "mediumDashDotDot": 11, "slantDashDot": 12, "thick": 13,
}
_BORDER_FROM_UNIVER = {v: k for k, v in _BORDER_TO_UNIVER.items()}
_HALIGN_TO_UNIVER = {
    "left": 1, "center": 2, "centerContinuous": 2, "right": 3,
    "justify": 4, "distributed": 6,
}
_HALIGN_FROM_UNIVER = {1: "left", 2: "center", 3: "right", 4: "justify", 5: "justify", 6: "distributed"}
_VALIGN_TO_UNIVER = {"top": 1, "center": 2, "bottom": 3, "justify": 2, "distributed": 2}
_VALIGN_FROM_UNIVER = {1: "top", 2: "center", 3: "bottom"}
T_STRING, T_NUMBER, T_BOOLEAN, T_FORCE_STRING = 1, 2, 3, 4

_COLOR_RE = re.compile(r"^#?([0-9a-fA-F]{6}|[0-9a-fA-F]{3})$")
_RGB_RE = re.compile(r"^rgba?\(\s*(\d+)\s*,\s*(\d+)\s*,\s*(\d+)")


class WorkbookRefused(HTTPException):
    """A save the editor must not do; ``status_code`` says why."""


# ─── what a snapshot may be (checked before any openpyxl work) ──────
# The characters an .xlsx can't hold (openpyxl refuses them mid-save).
_ILLEGAL_CHARS = re.compile(r"[\x00-\x08\x0b\x0c\x0e-\x1f]")


def _bad(detail: str, status: int = 400) -> WorkbookRefused:
    return WorkbookRefused(status_code=status, detail=detail)


def _check_text(value: Any, what: str, limit: int = MAX_TEXT) -> None:
    if value is None:
        return
    if not isinstance(value, str):
        raise _bad(f"{what} must be text")
    if len(value) > limit:
        raise _bad(f"{what} is longer than {limit} characters", 413)
    if _ILLEGAL_CHARS.search(value):
        raise _bad(f"{what} contains control characters a spreadsheet can't store")


def _check_number(value: Any, what: str, low: float, high: float) -> None:
    if value is None:
        return
    if isinstance(value, bool) or not isinstance(value, (int, float)):
        raise _bad(f"{what} must be a number")
    try:
        f = float(value)
    except OverflowError:
        raise _bad(f"{what} is out of range") from None
    if not math.isfinite(f) or f < low or f > high:
        raise _bad(f"{what} is out of range")


def _check_int(value: Any, what: str, low: int, high: int) -> int:
    if isinstance(value, bool) or not isinstance(value, int):
        raise _bad(f"{what} must be a whole number")
    if value < low or value > high:
        raise _bad(f"{what} is outside the sheet")
    return value


def _check_flag(value: Any, what: str) -> None:
    """An enum or 0/1 flag: a plain whole number (or true/false), nothing else."""
    if value is None or isinstance(value, bool):
        return
    if not isinstance(value, int) or not -1000 <= value <= 1000:
        raise _bad(f"{what} must be a whole number")


def _check_color(value: Any, what: str) -> None:
    if value is None:
        return
    if not isinstance(value, dict):
        raise _bad(f"{what} must be an object")
    rgb = value.get("rgb")
    if rgb is not None and (not isinstance(rgb, str) or len(rgb) > 64):
        raise _bad(f"{what} must be a colour")


def _check_style(style: Any, what: str) -> None:
    if style is None or isinstance(style, str):
        if isinstance(style, str) and len(style) > 64:
            raise _bad(f"{what}: style id too long")
        return
    if not isinstance(style, dict):
        raise _bad(f"{what}: a style must be an object")
    _check_text(style.get("ff"), f"{what}: font name", 255)
    _check_number(style.get("fs"), f"{what}: font size", 0, 409)
    n = style.get("n")
    if n is not None:
        if not isinstance(n, dict):
            raise _bad(f"{what}: number format must be an object")
        _check_text(n.get("pattern"), f"{what}: number format", 255)
    tr = style.get("tr")
    if tr is not None:
        if not isinstance(tr, dict):
            raise _bad(f"{what}: rotation must be an object")
        _check_number(tr.get("a"), f"{what}: rotation", -360, 360)
        _check_flag(tr.get("v"), f"{what}: rotation")
    for key in ("bl", "it", "ht", "vt", "tb", "stf", "td", "va"):
        _check_flag(style.get(key), f"{what}: {key}")
    for key in ("ul", "st", "ol", "bbl"):
        dec = style.get(key)
        if dec is None:
            continue
        if not isinstance(dec, dict):
            raise _bad(f"{what}: {key} must be an object")
        _check_flag(dec.get("s"), f"{what}: {key}")
    for key in ("cl", "bg"):
        _check_color(style.get(key), f"{what}: {key}")
    bd = style.get("bd")
    if bd is not None:
        if not isinstance(bd, dict):
            raise _bad(f"{what}: borders must be an object")
        for side in bd.values():
            if side is None:
                continue
            if not isinstance(side, dict):
                raise _bad(f"{what}: a border must be an object")
            _check_flag(side.get("s"), f"{what}: border style")
            _check_color(side.get("cl"), f"{what}: border colour")


def _check_cell(cell: Any, where: str, styles: dict) -> None:
    if cell is None:
        return
    if not isinstance(cell, dict):
        raise _bad(f"{where}: a cell must be an object")
    v = cell.get("v")
    if isinstance(v, str):
        _check_text(v, where)
    elif v is not None and not isinstance(v, bool):
        if not isinstance(v, (int, float)):
            raise _bad(f"{where}: a value must be text, a number or true/false")
        try:
            float(v)
        except OverflowError:
            raise _bad(f"{where}: the number is out of range") from None
    _check_text(cell.get("f"), f"{where} formula", 8192)
    p = cell.get("p")
    if p is not None:
        text = _rich_text(p)
        if text is None:
            raise _bad(f"{where}: unreadable rich text")
        _check_text(text, where)
    _check_style(cell.get("s"), where)


def _check_merges(raw: Any, name: str) -> None:
    """A list of non-overlapping ranges inside the window, at most
    ``MAX_MERGES`` of them covering at most ``MAX_MERGED_AREA`` cells.
    Overlaps are what made openpyxl rebuild huge merges over and over."""
    if raw is None:
        return
    if not isinstance(raw, list):
        raise _bad(f"{name}: merged cells must be a list")
    if len(raw) > MAX_MERGES:
        raise _bad(f"{name} has more than {MAX_MERGES} merged ranges", 413)
    spans = []
    area = 0
    for m in raw:
        if not isinstance(m, dict):
            raise _bad(f"{name}: a merged range must be an object")
        r0 = _check_int(m.get("startRow"), f"{name}: merge row", 0, WINDOW_ROWS - 1)
        r1 = _check_int(m.get("endRow"), f"{name}: merge row", 0, WINDOW_ROWS - 1)
        c0 = _check_int(m.get("startColumn"), f"{name}: merge column", 0, WINDOW_COLS - 1)
        c1 = _check_int(m.get("endColumn"), f"{name}: merge column", 0, WINDOW_COLS - 1)
        if r1 < r0 or c1 < c0:
            raise _bad(f"{name}: a merged range ends before it starts")
        area += (r1 - r0 + 1) * (c1 - c0 + 1)
        if area > MAX_MERGED_AREA:
            raise _bad(f"{name}: merged ranges cover more than {MAX_MERGED_AREA} cells", 413)
        spans.append((r0, r1, c0, c1))
    if len(spans) < 2:
        return
    # Overlap check in time proportional to the (bounded) merged area.
    grid = bytearray(WINDOW_ROWS * WINDOW_COLS)
    for r0, r1, c0, c1 in spans:
        width = c1 - c0 + 1
        for r in range(r0, r1 + 1):
            at = r * WINDOW_COLS + c0
            if any(grid[at:at + width]):
                raise _bad(f"{name}: merged ranges overlap")
            grid[at:at + width] = b"\x01" * width


def _check_dims(raw: Any, key: str, limit: int, high: int, what: str) -> None:
    if raw is None:
        return
    if not isinstance(raw, dict):
        raise _bad(f"{what} must be an object")
    if len(raw) > limit:
        raise _bad(f"{what}: too many entries", 413)
    for v in raw.values():
        if v is None:
            continue
        if not isinstance(v, dict):
            raise _bad(f"{what}: an entry must be an object")
        _check_number(v.get(key), what, 0, high)


def validate_snapshot(snapshot: Any) -> None:
    """Refuse (400/413) a snapshot that isn't one the editor could have
    sent, or that is bigger than this server will build: shapes, types,
    control characters, numbers out of range, overlapping merges and the
    size limits above. Runs before any openpyxl work."""
    if not isinstance(snapshot, dict):
        raise _bad("the workbook must be an object")
    order = snapshot.get("sheetOrder")
    sheets = snapshot.get("sheets")
    if not isinstance(order, list) or not order or not isinstance(sheets, dict):
        raise _bad("the workbook has no sheets")
    if len(order) > MAX_SHEETS:
        raise _bad(f"a workbook here can have at most {MAX_SHEETS} sheets", 413)
    if len(set(map(str, order))) != len(order):
        raise _bad("a sheet is listed twice")
    styles = snapshot.get("styles")
    if styles is not None:
        if not isinstance(styles, dict):
            raise _bad("styles must be an object")
        if len(styles) > MAX_STYLES:
            raise _bad("too many styles", 413)
        for sid, st in styles.items():
            if st is not None and not isinstance(st, dict):
                raise _bad("a style must be an object")
            _check_style(st, f"style {str(sid)[:20]}")
    cells = 0
    for sid in order:
        if not isinstance(sid, str):
            raise _bad("sheet ids must be text")
        sheet = sheets.get(sid)
        if not isinstance(sheet, dict):
            raise _bad(f"sheet {sid[:20]!r} is missing")
        name = sheet.get("name")
        _check_text(name, "A sheet name", 200)
        label = (name or sid)[:40]
        cd = sheet.get("cellData")
        if cd is not None:
            if not isinstance(cd, dict):
                raise _bad(f"{label}: cellData must be an object")
            for row in cd.values():
                if row is None:
                    continue
                if not isinstance(row, dict):
                    raise _bad(f"{label}: a row must be an object")
                cells += len(row)
                if cells > MAX_SAVE_CELLS:
                    raise _bad(
                        f"The workbook has more than {MAX_SAVE_CELLS:,} cells, which is more "
                        "than the editor saves. Nothing was changed.", 413)
                for cell in row.values():
                    _check_cell(cell, label, styles or {})
        _check_merges(sheet.get("mergeData"), label)
        _check_dims(sheet.get("columnData"), "w", WINDOW_COLS * 2, MAX_COL_PX, f"{label}: column widths")
        _check_dims(sheet.get("rowData"), "h", WINDOW_ROWS * 2, MAX_ROW_PX, f"{label}: row heights")
        freeze = sheet.get("freeze")
        if freeze is not None:
            if not isinstance(freeze, dict):
                raise _bad(f"{label}: freeze must be an object")
            for k in ("xSplit", "ySplit", "startRow", "startColumn"):
                if freeze.get(k) is not None:
                    _check_int(freeze.get(k), f"{label}: freeze", -1, WINDOW_ROWS)
        tab = sheet.get("tabColor")
        if tab is not None and (not isinstance(tab, str) or len(tab) > 32):
            raise _bad(f"{label}: tab colour must be a colour")
        origin = sheet.get("origin")
        if origin is not None and (not isinstance(origin, str) or len(origin) > 16):
            raise _bad(f"{label}: origin must be a sheet id")


# ─── small conversions ──────────────────────────────────────────────
def norm_color(value: Any) -> Optional[str]:
    """Any colour spelling Univer or openpyxl uses → ``#RRGGBB`` (or None)."""
    if value is None:
        return None
    if isinstance(value, dict):
        value = value.get("rgb")
        if value is None:
            return None
    s = str(value).strip()
    m = _COLOR_RE.match(s)
    if m:
        h = m.group(1)
        if len(h) == 3:
            h = "".join(ch * 2 for ch in h)
        return "#" + h.upper()
    m = _RGB_RE.match(s)
    if m:
        return "#" + "".join(f"{max(0, min(255, int(x))):02X}" for x in m.groups())
    if len(s) == 8 and re.fullmatch(r"[0-9a-fA-F]{8}", s):  # ARGB
        return "#" + s[2:].upper()
    return None


def _openpyxl_hex(color: Any) -> Optional[str]:
    """An openpyxl ``Color`` as ``#RRGGBB``; None for theme/auto colours,
    which have no RGB of their own (they are left untouched on save)."""
    if color is None:
        return None
    kind = getattr(color, "type", None)
    if kind == "rgb":
        rgb = color.rgb
        if isinstance(rgb, str) and len(rgb) in (6, 8):
            return "#" + rgb[-6:].upper()
        return None
    if kind == "indexed":
        from openpyxl.styles.colors import COLOR_INDEX

        idx = color.indexed
        if isinstance(idx, int) and 0 <= idx < 64 and idx < len(COLOR_INDEX):
            return "#" + COLOR_INDEX[idx][-6:].upper()
    return None


def width_to_px(chars: float) -> int:
    """Excel column width (characters of the default font) → pixels."""
    return int(round(chars * 7 + 5))


def px_to_width(px: float) -> float:
    return round(max(0.0, (float(px) - 5) / 7), 2)


def pt_to_px(pt: float) -> int:
    return int(round(float(pt) * 4 / 3))


def px_to_pt(px: float) -> float:
    return round(float(px) * 0.75, 2)


def _excel_serial(value: Any, epoch: Any) -> Optional[float]:
    from openpyxl.utils.datetime import to_excel

    try:
        if isinstance(value, _dt.timedelta):
            return value.total_seconds() / 86400.0
        return float(to_excel(value, epoch))
    except Exception:  # noqa: BLE001 — a date openpyxl itself can't place
        return None


# ─── value models: what a cell holds, comparable across both sides ──
def value_model(raw: Any, epoch: Any = None) -> Optional[tuple]:
    """An openpyxl value → ``('f', text) | ('n', float) | ('b', bool) |
    ('s', text)`` or None for empty."""
    if raw is None:
        return None
    try:
        from openpyxl.worksheet.formula import ArrayFormula

        if isinstance(raw, ArrayFormula):
            return ("f", raw.text or "")
    except ImportError:  # pragma: no cover
        pass
    if isinstance(raw, bool):
        return ("b", raw)
    if isinstance(raw, (int, float)):
        if isinstance(raw, float) and not math.isfinite(raw):
            return ("s", str(raw))
        return ("n", float(raw))
    if isinstance(raw, (_dt.datetime, _dt.date, _dt.time, _dt.timedelta)):
        serial = _excel_serial(raw, epoch)
        return ("n", serial) if serial is not None else ("s", str(raw))
    s = str(raw)
    if s == "":
        return None
    if s.startswith("="):
        return ("f", s)
    return ("s", s)


def _rich_text(p: Any) -> Optional[str]:
    try:
        stream = p["body"]["dataStream"]
    except (KeyError, TypeError):
        return None
    if not isinstance(stream, str):
        return None
    return stream.replace("\r\n", "\n").rstrip("\n").rstrip("\r")


def univer_value_model(cell: dict) -> Optional[tuple]:
    """A Univer ``ICellData`` → the same model as :func:`value_model`."""
    f = cell.get("f")
    if isinstance(f, str) and f.strip():
        return ("f", f if f.startswith("=") else "=" + f)
    v = cell.get("v")
    t = cell.get("t")
    if v is None and cell.get("p") is not None:
        v = _rich_text(cell.get("p"))
        t = t or T_STRING
    if v is None:
        return None
    if t == T_BOOLEAN or isinstance(v, bool):
        if isinstance(v, str):
            return ("b", v.strip().upper() == "TRUE")
        return ("b", bool(v))
    if t == T_FORCE_STRING:
        s = str(v)
        return ("s", s) if s != "" else None
    if isinstance(v, (int, float)) and not isinstance(v, bool):
        if not math.isfinite(float(v)):
            return ("s", str(v))
        return ("n", float(v))
    s = str(v)
    if s == "":
        return None
    if t == T_NUMBER:
        try:
            return ("n", float(s))
        except ValueError:
            return ("s", s)
    return ("s", s)


def _same_value(a: Optional[tuple], b: Optional[tuple]) -> bool:
    if a is None or b is None:
        return a is b
    if a[0] != b[0]:
        return False
    if a[0] == "n":
        x, y = a[1], b[1]
        return x == y or abs(x - y) <= 1e-9 * max(1.0, abs(x), abs(y))
    return a[1] == b[1]


def _to_univer_cell(model: Optional[tuple]) -> dict:
    if model is None:
        return {}
    kind, val = model
    if kind == "f":
        return {"f": val}
    if kind == "n":
        num = int(val) if float(val).is_integer() and abs(val) < 2**53 else val
        return {"v": num, "t": T_NUMBER}
    if kind == "b":
        return {"v": 1 if val else 0, "t": T_BOOLEAN}
    return {"v": val, "t": T_STRING}


# ─── style models ───────────────────────────────────────────────────
def style_model(cell: Any, default_font: Any) -> dict:
    """The parts of an openpyxl cell's style the editor shows, as a
    Univer ``IStyleData`` holding only what differs from the default."""
    out: dict[str, Any] = {}
    font = cell.font
    if font is not None:
        if font.name and default_font is not None and font.name != default_font.name:
            out["ff"] = font.name
        if font.sz and default_font is not None and float(font.sz) != float(default_font.sz or 11):
            out["fs"] = float(font.sz)
        if font.b:
            out["bl"] = 1
        if font.i:
            out["it"] = 1
        if font.u and font.u != "none":
            out["ul"] = {"s": 1}
        if font.strike:
            out["st"] = {"s": 1}
        cl = _openpyxl_hex(font.color)
        if cl and cl != "#000000":
            out["cl"] = {"rgb": cl}
    fill = cell.fill
    if fill is not None and getattr(fill, "fill_type", None) == "solid":
        bg = _openpyxl_hex(fill.fgColor)
        if bg:
            out["bg"] = {"rgb": bg}
    al = cell.alignment
    if al is not None:
        if al.horizontal in _HALIGN_TO_UNIVER:
            out["ht"] = _HALIGN_TO_UNIVER[al.horizontal]
        if al.vertical in _VALIGN_TO_UNIVER:
            out["vt"] = _VALIGN_TO_UNIVER[al.vertical]
        if al.wrap_text:
            out["tb"] = 3
        if al.shrink_to_fit:
            out["stf"] = 1
        # Excel: 1-90 turns the text up (anticlockwise), 91-180 down by
        # (n - 90), 255 stacks it. Univer: negative turns it up.
        rot = int(al.text_rotation or 0)
        if rot == 255:
            out["tr"] = {"a": 0, "v": 1}
        elif 0 < rot <= 90:
            out["tr"] = {"a": -rot}
        elif 90 < rot <= 180:
            out["tr"] = {"a": rot - 90}
    bd = cell.border
    if bd is not None:
        sides = {}
        for key, side in (("t", bd.top), ("b", bd.bottom), ("l", bd.left), ("r", bd.right)):
            if side is not None and side.style in _BORDER_TO_UNIVER:
                sides[key] = {
                    "s": _BORDER_TO_UNIVER[side.style],
                    "cl": {"rgb": _openpyxl_hex(side.color) or "#000000"},
                }
        if sides:
            out["bd"] = sides
    fmt = cell.number_format
    if fmt and fmt != "General":
        out["n"] = {"pattern": fmt}
    return out


def norm_style(style: Any) -> dict:
    """A Univer style (as the editor sends it back) reduced to the keys and
    spellings :func:`style_model` produces, so "nothing changed" compares
    equal however Univer chose to spell it."""
    if not isinstance(style, dict):
        return {}
    out: dict[str, Any] = {}
    ff = style.get("ff")
    if isinstance(ff, str) and ff.strip():
        out["ff"] = ff.strip()
    fs = style.get("fs")
    if isinstance(fs, (int, float)) and not isinstance(fs, bool) and fs > 0:
        out["fs"] = float(fs)
    for key in ("bl", "it"):
        if style.get(key) in (1, True):
            out[key] = 1
    for key in ("ul", "st"):
        dec = style.get(key)
        if isinstance(dec, dict) and dec.get("s") in (1, True):
            out[key] = {"s": 1}
    for key in ("cl", "bg"):
        c = norm_color(style.get(key))
        if c and not (key == "cl" and c == "#000000"):
            out[key] = {"rgb": c}
    ht = style.get("ht")
    if ht in _HALIGN_FROM_UNIVER:
        out["ht"] = 4 if ht == 5 else ht
    vt = style.get("vt")
    if vt in _VALIGN_FROM_UNIVER:
        out["vt"] = vt
    if style.get("tb") == 3:
        out["tb"] = 3
    if style.get("stf") in (1, True):
        out["stf"] = 1
    tr = style.get("tr")
    if isinstance(tr, dict):
        if tr.get("v") in (1, True):
            out["tr"] = {"a": 0, "v": 1}
        else:
            try:
                a = int(round(float(tr.get("a") or 0)))
            except (TypeError, ValueError):
                a = 0
            a = max(-90, min(90, a))
            if a:
                out["tr"] = {"a": a}
    bd = style.get("bd")
    if isinstance(bd, dict):
        sides = {}
        for key in ("t", "b", "l", "r"):
            side = bd.get(key)
            if isinstance(side, dict) and side.get("s") in _BORDER_FROM_UNIVER:
                sides[key] = {"s": side["s"], "cl": {"rgb": norm_color(side.get("cl")) or "#000000"}}
        if sides:
            out["bd"] = sides
    n = style.get("n")
    if isinstance(n, dict) and isinstance(n.get("pattern"), str) and n["pattern"] not in ("", "General"):
        out["n"] = {"pattern": n["pattern"]}
    return out


def _apply_style(cell: Any, want: dict, have: dict, default_font: Any) -> None:
    """Set on ``cell`` only the style attributes where ``want`` differs
    from ``have`` (the cell's current :func:`style_model`). Everything
    else on its font, fill, alignment and border objects stays."""
    from openpyxl.styles import Alignment, Border, PatternFill, Side
    from openpyxl.styles.colors import Color

    def changed(*keys: str) -> bool:
        return any(want.get(k) != have.get(k) for k in keys)

    if changed("ff", "fs", "bl", "it", "ul", "st", "cl"):
        font = copy(cell.font)
        if changed("ff"):
            font.name = want.get("ff") or (default_font.name if default_font is not None else "Calibri")
        if changed("fs"):
            font.sz = want.get("fs") or (default_font.sz if default_font is not None else 11)
        if changed("bl"):
            font.b = bool(want.get("bl"))
        if changed("it"):
            font.i = bool(want.get("it"))
        if changed("ul"):
            font.u = "single" if want.get("ul") else None
        if changed("st"):
            font.strike = bool(want.get("st"))
        if changed("cl"):
            c = (want.get("cl") or {}).get("rgb")
            font.color = Color(rgb="FF" + c[1:]) if c else None
        cell.font = font
    if changed("bg"):
        c = (want.get("bg") or {}).get("rgb")
        cell.fill = PatternFill(fill_type="solid", fgColor="FF" + c[1:], bgColor="FF" + c[1:]) \
            if c else PatternFill(fill_type=None)
    if changed("ht", "vt", "tb", "stf", "tr"):
        al = copy(cell.alignment) if cell.alignment is not None else Alignment()
        kw = {
            "horizontal": al.horizontal, "vertical": al.vertical,
            "wrap_text": al.wrap_text, "shrink_to_fit": al.shrink_to_fit,
            "text_rotation": al.text_rotation, "indent": al.indent,
        }
        if changed("ht"):
            kw["horizontal"] = _HALIGN_FROM_UNIVER.get(want.get("ht"))
        if changed("vt"):
            kw["vertical"] = _VALIGN_FROM_UNIVER.get(want.get("vt"))
        if changed("tb"):
            kw["wrap_text"] = True if want.get("tb") == 3 else None
        if changed("stf"):
            kw["shrink_to_fit"] = True if want.get("stf") else None
        if changed("tr"):
            tr = want.get("tr") or {}
            if tr.get("v"):
                kw["text_rotation"] = 255
            else:
                a = int(tr.get("a") or 0)
                kw["text_rotation"] = -a if a <= 0 else 90 + a
        cell.alignment = Alignment(**kw)
    if changed("bd"):
        cur = cell.border
        wb_sides = want.get("bd") or {}
        hb_sides = have.get("bd") or {}

        def side_for(key: str, existing: Any) -> Any:
            if wb_sides.get(key) == hb_sides.get(key):
                return existing
            spec = wb_sides.get(key)
            if not spec:
                return Side()
            rgb = (spec.get("cl") or {}).get("rgb") or "#000000"
            return Side(style=_BORDER_FROM_UNIVER[spec["s"]], color="FF" + rgb[1:])

        cell.border = Border(
            left=side_for("l", cur.left), right=side_for("r", cur.right),
            top=side_for("t", cur.top), bottom=side_for("b", cur.bottom),
            diagonal=cur.diagonal, diagonal_direction=cur.diagonal_direction,
            diagonalUp=cur.diagonalUp, diagonalDown=cur.diagonalDown,
            outline=cur.outline, vertical=cur.vertical, horizontal=cur.horizontal,
        )
    if changed("n"):
        cell.number_format = (want.get("n") or {}).get("pattern") or "General"


# ─── read ───────────────────────────────────────────────────────────
class _Styles:
    """Interns style dicts into the workbook's shared ``styles`` table."""

    def __init__(self) -> None:
        self.table: dict[str, dict] = {}
        self._ids: dict[str, str] = {}

    def ref(self, style: dict) -> Optional[str]:
        if not style:
            return None
        import json

        key = json.dumps(style, sort_keys=True)
        sid = self._ids.get(key)
        if sid is None:
            sid = f"st{len(self._ids) + 1}"
            self._ids[key] = sid
            self.table[sid] = style
        return sid


def _no_freeze() -> dict:
    return {"xSplit": 0, "ySplit": 0, "startRow": -1, "startColumn": -1}


def _freeze_of(ws: Any) -> dict:
    from openpyxl.utils.cell import coordinate_to_tuple

    fp = ws.freeze_panes
    if not fp:
        return _no_freeze()
    try:
        row, col = coordinate_to_tuple(fp)
    except Exception:  # noqa: BLE001
        return _no_freeze()
    rows, cols = row - 1, col - 1
    if rows <= 0 and cols <= 0:
        return _no_freeze()
    return {"xSplit": max(cols, 0), "ySplit": max(rows, 0), "startRow": max(rows, 0), "startColumn": max(cols, 0)}


def _column_model(ws: Any, n_cols: int) -> dict[int, dict]:
    out: dict[int, dict] = {}
    for dim in ws.column_dimensions.values():
        lo, hi = int(dim.min or 0), int(dim.max or 0)
        if lo < 1:
            continue
        width = dim.width if dim.customWidth or dim.width else None
        for c in range(lo, min(hi, n_cols) + 1):
            m: dict[str, Any] = {}
            if width:
                m["w"] = width_to_px(width)
            if dim.hidden:
                m["hd"] = 1
            if m:
                out[c - 1] = m
    return out


def _row_model(ws: Any, n_rows: int) -> dict[int, dict]:
    out: dict[int, dict] = {}
    for r, dim in ws.row_dimensions.items():
        if not isinstance(r, int) or r < 1 or r > n_rows:
            continue
        m: dict[str, Any] = {}
        if dim.height is not None and dim.customHeight is not False:
            m["h"] = pt_to_px(dim.height)
        if dim.hidden:
            m["hd"] = 1
        if m:
            out[r - 1] = m
    return out


def _merges_of(ws: Any, n_rows: int, n_cols: int) -> list[dict]:
    out = []
    for rng in ws.merged_cells.ranges:
        if rng.min_row > n_rows or rng.min_col > n_cols:
            continue
        out.append({
            "startRow": rng.min_row - 1, "endRow": min(rng.max_row, n_rows) - 1,
            "startColumn": rng.min_col - 1, "endColumn": min(rng.max_col, n_cols) - 1,
        })
    return out


def _sheet_extent(ws: Any) -> tuple[int, int]:
    """The rows/columns that hold anything (value or style), from the cells
    openpyxl actually created rather than ``max_row``, which a stray
    format far down the sheet inflates."""
    rows = cols = 0
    for (r, c), cell in ws._cells.items():
        if cell.value is not None or cell.has_style:
            rows = max(rows, r)
            cols = max(cols, c)
    for rng in ws.merged_cells.ranges:
        rows = max(rows, rng.max_row)
        cols = max(cols, rng.max_col)
    return rows, cols


def _read_xlsx(target: Path, wb: Any) -> dict:
    from openpyxl.cell.cell import MergedCell

    styles = _Styles()
    default_font = wb._fonts[0] if getattr(wb, "_fonts", None) else None
    sheets: dict[str, dict] = {}
    order: list[str] = []
    truncated: list[str] = []
    budget = MAX_OPEN_CELLS
    over = False
    for idx, ws in enumerate(wb.worksheets):
        sid = f"s{idx}"
        order.append(sid)
        used_rows, used_cols = _sheet_extent(ws)
        if used_rows > WINDOW_ROWS or used_cols > WINDOW_COLS:
            truncated.append(ws.title)
        n_rows, n_cols = min(used_rows, WINDOW_ROWS), min(used_cols, WINDOW_COLS)
        cell_data: dict[int, dict[int, dict]] = {}
        for (r, c), cell in ws._cells.items():
            if r > n_rows or c > n_cols or isinstance(cell, MergedCell):
                continue
            if budget <= 0:
                over = True
                break
            out = _to_univer_cell(value_model(cell.value, wb.epoch))
            sref = styles.ref(style_model(cell, default_font)) if cell.has_style else None
            if sref:
                out["s"] = sref
            if out:
                cell_data.setdefault(r - 1, {})[c - 1] = out
                budget -= 1
        tab = _openpyxl_hex(ws.sheet_properties.tabColor) if ws.sheet_properties.tabColor else None
        sheet: dict[str, Any] = {
            "id": sid,
            "name": ws.title,
            "hidden": 1 if ws.sheet_state in ("hidden", "veryHidden") else 0,
            "rowCount": min(WINDOW_ROWS, max(1000, n_rows + 100)),
            "columnCount": min(WINDOW_COLS, max(26, n_cols + 5)),
            "cellData": cell_data,
            "mergeData": _merges_of(ws, n_rows, n_cols),
            "columnData": _column_model(ws, WINDOW_COLS),
            "rowData": _row_model(ws, WINDOW_ROWS),
            "freeze": _freeze_of(ws),
            "showGridlines": 0 if ws.sheet_view.showGridLines is False else 1,
        }
        if tab:
            sheet["tabColor"] = tab
        sheets[sid] = sheet
    workbook: dict[str, Any] = {
        "id": "domovoi-sheet",
        "name": target.stem,
        "appVersion": "domovoi",
        "locale": "enUS",
        "styles": styles.table,
        "sheetOrder": order,
        "sheets": sheets,
    }
    if default_font is not None and default_font.name:
        workbook["defaultStyle"] = {"ff": default_font.name, "fs": float(default_font.sz or 11)}
    return {"workbook": workbook, "truncated": truncated, "over_budget": over}


def _csv_cell(text: str) -> dict:
    if text == "":
        return {}
    if text.startswith("="):
        return {"f": text}
    # A number only when it reads back exactly the same (so "007", "1.50"
    # and "1e3" stay the text they are).
    try:
        num = float(text)
    except ValueError:
        return {"v": text, "t": T_STRING}
    if math.isfinite(num):
        as_num = int(num) if num.is_integer() and "." not in text and "e" not in text.lower() else num
        if str(as_num) == text:
            return {"v": as_num, "t": T_NUMBER}
    return {"v": text, "t": T_STRING}


def csv_text(cell: Optional[dict]) -> str:
    """What a Univer cell is as CSV text: the formula, or the value."""
    m = univer_value_model(cell or {})
    if m is None:
        return ""
    kind, val = m
    if kind == "n":
        return str(int(val)) if float(val).is_integer() and abs(val) < 2**53 else repr(val)
    if kind == "b":
        return "TRUE" if val else "FALSE"
    return str(val)


def read_csv(target: Path, max_rows: int, max_cols: int) -> dict:
    cell_data: dict[int, dict[int, dict]] = {}
    n_rows = n_cols = 0
    truncated = False
    budget = MAX_OPEN_CELLS
    over = False
    with target.open("r", encoding="utf-8-sig", errors="replace", newline="") as fh:
        for r, row in enumerate(csv.reader(fh)):
            if r >= max_rows:
                truncated = True
                break
            if budget <= 0:
                over = True
                break
            budget -= len(row)
            if len(row) > max_cols:
                truncated = True
            for c, text in enumerate(row[:max_cols]):
                out = _csv_cell(text)
                if out:
                    cell_data.setdefault(r, {})[c] = out
                    n_cols = max(n_cols, c + 1)
            n_rows = r + 1
    sheet = {
        "id": "s0",
        "name": target.stem[:31] or "Sheet1",
        "hidden": 0,
        "rowCount": max(max_rows, 100) if truncated else min(max_rows, max(1000, n_rows + 100)),
        "columnCount": min(max_cols, max(26, n_cols + 5)),
        "cellData": cell_data,
        "mergeData": [],
        "columnData": {},
        "rowData": {},
        "freeze": _no_freeze(),
    }
    return {
        "workbook": {
            "id": "domovoi-sheet", "name": target.stem, "appVersion": "domovoi",
            "locale": "enUS", "styles": {}, "sheetOrder": ["s0"], "sheets": {"s0": sheet},
        },
        "truncated": [sheet["name"]] if truncated else [],
        "over_budget": over,
    }


# openpyxl's own word for "this part of the file is dropped when I save it"
# (sparklines, x14 validation and conditional formatting, slicers, ...).
# ANY warning openpyxl gives while loading means it may not write the file
# back as it found it (round 2 proved a dynamic Print_Area lost under the
# wording "Print area cannot be set to ..."), so every UserWarning makes the
# file read-only, except these, which change nothing a save would write
# differently.
_HARMLESS_WARNINGS = (
    "Workbook contains no default style",
    "Workbook contains no stylesheet",
    "Title is more than 31 characters",
)


def load_xlsx(target: Path, **kwargs: Any) -> tuple[Any, list[str]]:
    """``load_workbook`` plus what openpyxl said it could not keep. A
    workbook that loses anything on a load-and-save must not be saved by
    the editor; the caller opens it read-only (or refuses the save)."""
    import warnings

    from openpyxl import load_workbook

    with warnings.catch_warnings(record=True) as caught:
        warnings.simplefilter("always")
        wb = load_workbook(target, **kwargs)
    lost = sorted({
        str(w.message) for w in caught
        if issubclass(w.category, UserWarning)
        and not str(w.message).startswith(_HARMLESS_WARNINGS)
    })
    return wb, lost


def read_xlsx(target: Path) -> dict:
    wb, lost = load_xlsx(target, data_only=False, rich_text=False)
    try:
        model = _read_xlsx(target, wb)
        model["lost"] = lost
        return model
    finally:
        wb.close()


def read_xlsx_values_only(target: Path) -> dict:
    """A workbook too large to edit, for viewing: values and formulas of
    every sheet through openpyxl's streaming reader, no formatting."""
    from openpyxl import load_workbook

    wb = load_workbook(target, data_only=False, read_only=True)
    try:
        sheets: dict[str, dict] = {}
        order: list[str] = []
        truncated: list[str] = []
        # One budget for the whole workbook: the read stops when it is spent
        # (the rest of the file is simply not shown), so neither the read nor
        # the answer grows with the number of sheets.
        budget = MAX_OPEN_CELLS
        for idx, ws in enumerate(wb.worksheets):
            sid = f"s{idx}"
            order.append(sid)
            cell_data: dict[int, dict[int, dict]] = {}
            n_rows = n_cols = 0
            rows = ws.iter_rows(max_row=WINDOW_ROWS + 1, max_col=WINDOW_COLS) if budget > 0 else ()
            if budget <= 0:
                truncated.append(ws.title)
            for r, row in enumerate(rows):
                if r >= WINDOW_ROWS or budget <= 0:
                    truncated.append(ws.title)
                    break
                for c, cell in enumerate(row):
                    out = _to_univer_cell(value_model(getattr(cell, "value", None), wb.epoch))
                    if out:
                        cell_data.setdefault(r, {})[c] = out
                        n_rows, n_cols = max(n_rows, r + 1), max(n_cols, c + 1)
                        budget -= 1
            sheets[sid] = {
                "id": sid, "name": ws.title, "hidden": 0,
                "rowCount": min(WINDOW_ROWS, max(1000, n_rows + 100)),
                "columnCount": min(WINDOW_COLS, max(26, n_cols + 5)),
                "cellData": cell_data, "freeze": _no_freeze(),
            }
        return {
            "workbook": {
                "id": "domovoi-sheet", "name": target.stem, "appVersion": "domovoi",
                "locale": "enUS", "styles": {}, "sheetOrder": order, "sheets": sheets,
            },
            "truncated": truncated,
            "over_budget": budget <= 0,
        }
    finally:
        wb.close()


# ─── write ──────────────────────────────────────────────────────────
def _int_key(k: Any) -> Optional[int]:
    try:
        n = int(k)
    except (TypeError, ValueError):
        return None
    return n if n >= 0 else None


def _resolve_style(raw: Any, table: dict) -> dict:
    if isinstance(raw, str):
        raw = table.get(raw)
    return norm_style(raw) if isinstance(raw, dict) else {}


def snapshot_cells(sheet: dict, styles: dict) -> dict[tuple[int, int], tuple[Optional[tuple], dict]]:
    """``{(row, col): (value_model, style)}`` for every cell the snapshot
    holds (0-based). Refuses (413) a cell outside the editor's window."""
    out: dict[tuple[int, int], tuple[Optional[tuple], dict]] = {}
    cell_data = sheet.get("cellData") or {}
    if not isinstance(cell_data, dict):
        raise WorkbookRefused(status_code=400, detail="cellData must be an object")
    for rk, row in cell_data.items():
        r = _int_key(rk)
        if r is None or not isinstance(row, dict):
            continue
        for ck, cell in row.items():
            c = _int_key(ck)
            if c is None or not isinstance(cell, dict):
                continue
            model = univer_value_model(cell)
            style = _resolve_style(cell.get("s"), styles)
            if model is None and not style:
                continue
            if r >= WINDOW_ROWS or c >= WINDOW_COLS:
                raise WorkbookRefused(
                    status_code=413,
                    detail=f"{sheet.get('name') or 'A sheet'} has data past row {WINDOW_ROWS} or "
                    f"column {WINDOW_COLS}, which the editor can't save. Nothing was changed.",
                )
            out[(r, c)] = (model, style)
    return out


def _snapshot_merges(sheet: dict) -> set[tuple[int, int, int, int]]:
    out = set()
    for m in sheet.get("mergeData") or []:
        try:
            r0, r1 = int(m["startRow"]), int(m["endRow"])
            c0, c1 = int(m["startColumn"]), int(m["endColumn"])
        except (KeyError, TypeError, ValueError):
            continue
        if 0 <= r0 <= r1 < WINDOW_ROWS and 0 <= c0 <= c1 < WINDOW_COLS and (r0, c0) != (r1, c1):
            out.add((r0, r1, c0, c1))
    return out


def _dims_model(raw: Any, key: str, limit: int) -> dict[int, dict]:
    out: dict[int, dict] = {}
    if not isinstance(raw, dict):
        return out
    for k, v in raw.items():
        i = _int_key(k)
        if i is None or i >= limit or not isinstance(v, dict):
            continue
        m: dict[str, Any] = {}
        size = v.get(key)
        if isinstance(size, (int, float)) and not isinstance(size, bool) and size > 0:
            m[key] = int(round(size))
        if v.get("hd") in (1, True):
            m["hd"] = 1
        if m:
            out[i] = m
    return out


def _explode_column_dims(ws: Any) -> None:
    """Give every grouped ``<col min max>`` its own per-column entry, so a
    width set on one column of the group never overlaps it (Excel calls an
    overlapping pair a damaged file)."""
    from openpyxl.utils import get_column_letter

    for key, dim in list(ws.column_dimensions.items()):
        lo, hi = int(dim.min or 0), int(dim.max or 0)
        if lo < 1 or hi <= lo:
            continue
        del ws.column_dimensions[key]
        for c in range(lo, hi + 1):
            letter = get_column_letter(c)
            new = copy(dim)
            new.index = letter
            new.min = new.max = c
            ws.column_dimensions[letter] = new


def _write_dims(ws: Any, have_cols: dict, want_cols: dict, have_rows: dict, want_rows: dict) -> None:
    from openpyxl.utils import get_column_letter

    col_keys = {k for k in set(have_cols) | set(want_cols) if have_cols.get(k) != want_cols.get(k)}
    # Univer reports a size it gave a column/row itself only when it differs
    # from its default; a column the file never sized and the editor left at
    # the default is not a change.
    if col_keys:
        _explode_column_dims(ws)
    for c in sorted(col_keys):
        want, have = want_cols.get(c, {}), have_cols.get(c, {})
        letter = get_column_letter(c + 1)
        existed = letter in ws.column_dimensions
        dim = ws.column_dimensions[letter]
        if not existed:
            dim.width = 0  # openpyxl's default would write a 13-character width
        if want.get("w") != have.get("w"):
            if want.get("w"):
                dim.width = px_to_width(want["w"])
            else:
                dim.width = 0  # no width of its own (0 is not written)
        if want.get("hd") != have.get("hd"):
            dim.hidden = bool(want.get("hd"))
    for r in sorted(k for k in set(have_rows) | set(want_rows) if have_rows.get(k) != want_rows.get(k)):
        want, have = want_rows.get(r, {}), have_rows.get(r, {})
        dim = ws.row_dimensions[r + 1]
        if want.get("h") != have.get("h"):
            if want.get("h"):
                dim.height = px_to_pt(want["h"])
            else:
                dim.height = None
        if want.get("hd") != have.get("hd"):
            dim.hidden = bool(want.get("hd"))


def _set_value(cell: Any, model: Optional[tuple]) -> None:
    if model is None:
        cell.value = None
        return
    kind, val = model
    if kind == "f":
        cell.value = val
    elif kind == "n":
        cell.value = int(val) if float(val).is_integer() and abs(val) < 2**53 else val
    elif kind == "b":
        cell.value = bool(val)
    else:
        cell.value = val
        if isinstance(val, str) and val.startswith("="):
            cell.data_type = "s"  # typed as text (Univer's force-string), not a formula


def _write_sheet(ws: Any, sheet: dict, styles: dict, wb: Any, fresh: bool) -> None:
    from openpyxl.cell.cell import MergedCell
    from openpyxl.utils import get_column_letter

    default_font = wb._fonts[0] if getattr(wb, "_fonts", None) else None
    want_cells = snapshot_cells(sheet, styles)

    used_rows, used_cols = (0, 0) if fresh else _sheet_extent(ws)
    n_rows, n_cols = min(used_rows, WINDOW_ROWS), min(used_cols, WINDOW_COLS)

    # Merges first: unmerge what went, so its cells can take values.
    have_merges = set() if fresh else {
        (m["startRow"], m["endRow"], m["startColumn"], m["endColumn"])
        for m in _merges_of(ws, n_rows, n_cols)
    }
    want_merges = _snapshot_merges(sheet)
    for (r0, r1, c0, c1) in have_merges - want_merges:
        ws.unmerge_cells(start_row=r0 + 1, end_row=r1 + 1, start_column=c0 + 1, end_column=c1 + 1)

    # Cells: everything the snapshot holds, plus every cell the file holds
    # inside the window (a cell missing from the snapshot was cleared).
    keys = set(want_cells)
    for (r, c), cell in list(ws._cells.items()):
        if r <= n_rows and c <= n_cols and (cell.value is not None or cell.has_style):
            keys.add((r - 1, c - 1))
    for (r, c) in sorted(keys):
        existing = ws._cells.get((r + 1, c + 1))
        if isinstance(existing, MergedCell):
            continue
        have_val = value_model(existing.value, wb.epoch) if existing is not None else None
        have_style = style_model(existing, default_font) if existing is not None and existing.has_style else {}
        want_val, want_style = want_cells.get((r, c), (None, {}))
        if default_font is not None and want_style:
            # Univer may spell out the workbook's default font on a cell it
            # restyled; the file says that by saying nothing.
            if want_style.get("ff") == default_font.name:
                want_style = {k: v for k, v in want_style.items() if k != "ff"}
            if want_style.get("fs") == float(default_font.sz or 11):
                want_style = {k: v for k, v in want_style.items() if k != "fs"}
        if existing is None and want_val is None and not want_style:
            continue
        cell = existing if existing is not None else ws.cell(row=r + 1, column=c + 1)
        if not _same_value(want_val, have_val):
            _set_value(cell, want_val)
        if want_style != have_style:
            _apply_style(cell, want_style, have_style, default_font)

    for (r0, r1, c0, c1) in sorted(want_merges - have_merges):
        ws.merge_cells(start_row=r0 + 1, end_row=r1 + 1, start_column=c0 + 1, end_column=c1 + 1)

    _write_dims(
        ws,
        {} if fresh else _column_model(ws, WINDOW_COLS),
        _dims_model(sheet.get("columnData"), "w", WINDOW_COLS),
        {} if fresh else _row_model(ws, WINDOW_ROWS),
        _dims_model(sheet.get("rowData"), "h", WINDOW_ROWS),
    )

    freeze = sheet.get("freeze") or {}
    try:
        rows = max(0, int(freeze.get("startRow", -1))) if int(freeze.get("ySplit", 0) or 0) > 0 else 0
        cols = max(0, int(freeze.get("startColumn", -1))) if int(freeze.get("xSplit", 0) or 0) > 0 else 0
    except (TypeError, ValueError):
        rows = cols = 0
    want_freeze = _freeze_of_counts(rows, cols)
    if fresh or want_freeze != _freeze_of(ws):
        ws.freeze_panes = f"{get_column_letter(cols + 1)}{rows + 1}" if (rows or cols) else None

    if "showGridlines" in sheet:
        show = sheet.get("showGridlines") not in (0, False)
        if (ws.sheet_view.showGridLines is not False) != show:
            ws.sheet_view.showGridLines = show

    tab = norm_color(sheet.get("tabColor"))
    have_tab = _openpyxl_hex(ws.sheet_properties.tabColor) if ws.sheet_properties.tabColor else None
    if tab != have_tab:
        ws.sheet_properties.tabColor = tab[1:] if tab else None


def _freeze_of_counts(rows: int, cols: int) -> dict:
    if rows <= 0 and cols <= 0:
        return _no_freeze()
    return {"xSplit": cols, "ySplit": rows, "startRow": rows, "startColumn": cols}


def _valid_title(name: Any, fallback: str) -> str:
    s = str(name or "").strip() or fallback
    s = re.sub(r"[\\/*?:\[\]]", "_", s)[:31]
    return s.strip("'") or fallback


def _mentions_sheet(text: str, title: str) -> bool:
    """``text`` (a formula, a defined name) refers to the sheet ``title``:
    ``Title!`` or ``'Ti''tle'!``, case-insensitively, as Excel reads it."""
    low, name = text.lower(), title.lower()
    if "'" + name.replace("'", "''") + "'!" in low:
        return True
    for m in re.finditer(re.escape(name) + "!", low):
        i = m.start()
        if i == 0 or not (low[i - 1].isalnum() or low[i - 1] in "_.'"):
            return True
    return False


def _file_references(wb: Any, skip: Any = None) -> list[str]:
    """Every formula-like text in the workbook that the editor doesn't show
    and so can't rewrite: defined names (workbook and sheet scope), data
    validation and conditional formatting formulas."""
    out: list[str] = []
    for dn in getattr(wb, "defined_names", {}).values():
        out.append(dn.attr_text or "")
    for ws in wb.worksheets:
        for dn in getattr(ws, "defined_names", {}).values():
            out.append(dn.attr_text or "")
        if ws is skip:
            continue
        for dv in ws.data_validations.dataValidation:
            out.extend(f for f in (dv.formula1, dv.formula2) if f)
        for cf in ws.conditional_formatting:
            for rule in cf.rules:
                out.extend(f for f in (rule.formula or []) if f)
    return out


def _snapshot_formulas(snapshot: dict) -> list[str]:
    out = []
    for sheet in (snapshot.get("sheets") or {}).values():
        for row in ((sheet or {}).get("cellData") or {}).values():
            for cell in (row or {}).values():
                f = (cell or {}).get("f") if isinstance(cell, dict) else None
                if isinstance(f, str) and f:
                    out.append(f)
    return out


def _refuse_breaking_renames(wb: Any, plan: list[tuple[Any, dict, bool]], gone: list[Any],
                             snapshot: dict) -> None:
    """Renaming or removing a sheet would leave named ranges, validation or
    formatting rules (and any formula Univer didn't rewrite) pointing at a
    sheet that isn't there. v1 refuses that save with a plain message
    rather than rewrite references it can't see."""
    renamed = [(ws, ws.title) for ws, sheet, fresh in plan
               if not fresh and str(sheet.get("name") or "").strip().lower() != ws.title.lower()]
    if not renamed and not gone:
        return
    texts = _file_references(wb)
    formulas = _snapshot_formulas(snapshot)
    for ws, old in renamed:
        if any(_mentions_sheet(t, old) for t in texts) or any(_mentions_sheet(f, old) for f in formulas):
            raise WorkbookRefused(
                status_code=409,
                detail=f"The sheet “{old}” is used by named ranges, rules or formulas in this "
                "file, so it can't be renamed here. Give it back its name to save, or rename it "
                "in a spreadsheet app.",
            )
    for ws in gone:
        if any(_mentions_sheet(t, ws.title) for t in _file_references(wb, skip=ws)):
            raise WorkbookRefused(
                status_code=409,
                detail=f"The sheet “{ws.title}” is used by named ranges or rules in this file, "
                "so it can't be deleted here. Delete it in a spreadsheet app.",
            )


def write_xlsx(wb: Any, snapshot: dict) -> None:
    """Apply ``snapshot`` (``{sheetOrder, sheets, styles}`` plus each
    sheet's ``origin``) to ``wb`` in place. The caller has run
    :func:`validate_snapshot`; this runs it again, cheaply, for any other
    caller."""
    validate_snapshot(snapshot)
    order = snapshot.get("sheetOrder")
    sheets = snapshot.get("sheets")
    styles = snapshot.get("styles") if isinstance(snapshot.get("styles"), dict) else {}
    origins = {f"s{i}": ws for i, ws in enumerate(wb.worksheets)}
    used: set[str] = set()
    plan: list[tuple[Any, dict, bool]] = []
    for sid in order:
        sheet = sheets.get(sid)
        if not isinstance(sheet, dict):
            raise WorkbookRefused(status_code=400, detail=f"sheet {sid!r} is missing")
        origin = sheet.get("origin")
        ws = origins.get(origin) if isinstance(origin, str) else None
        if ws is not None and origin in used:
            ws = None  # a duplicated sheet: a new copy, not the same one twice
        if ws is not None:
            used.add(origin)
        plan.append((ws, sheet, ws is None))
    _refuse_breaking_renames(wb, plan, [ws for o, ws in origins.items() if o not in used], snapshot)
    # Sheets that are gone, then new ones, then the order.
    for origin, ws in origins.items():
        if origin not in used:
            wb.remove(ws)
    kept = []
    for i, (ws, sheet, fresh) in enumerate(plan):
        if fresh:
            ws = wb.create_sheet(title=f"__domovoi_new_{i}")
        kept.append((ws, sheet, fresh))
    wb._sheets = [ws for ws, _, _ in kept]
    # Titles in two passes, so swapping two names never collides.
    for i, (ws, _, _) in enumerate(kept):
        ws.title = f"__domovoi_tmp_{i}"
    seen: set[str] = set()
    for i, (ws, sheet, _) in enumerate(kept):
        title = _valid_title(sheet.get("name"), f"Sheet{i + 1}")
        base, n = title, 2
        while title.lower() in seen:
            suffix = f" ({n})"
            title = base[: 31 - len(suffix)] + suffix
            n += 1
        seen.add(title.lower())
        ws.title = title
    for ws, sheet, fresh in kept:
        _write_sheet(ws, sheet, styles, wb, fresh)
        hidden = sheet.get("hidden") in (1, 2, True)
        if hidden and ws.sheet_state == "visible":
            ws.sheet_state = "hidden"
        elif not hidden and ws.sheet_state != "visible":
            ws.sheet_state = "visible"
    if all(ws.sheet_state != "visible" for ws, _, _ in kept):
        kept[0][0].sheet_state = "visible"
    first_visible = next(i for i, (ws, _, _) in enumerate(kept) if ws.sheet_state == "visible")
    wb.active = first_visible
    for ws, _, _ in kept:
        ws.sheet_view.tabSelected = ws is kept[first_visible][0]


def csv_rows(snapshot: dict) -> list[list[str]]:
    """The first sheet of a snapshot as CSV rows (values and formula text)."""
    order = snapshot.get("sheetOrder") or []
    sheets = snapshot.get("sheets") or {}
    if not order or not isinstance(sheets.get(order[0]), dict):
        raise WorkbookRefused(status_code=400, detail="the workbook has no sheets")
    cells = snapshot_cells(sheets[order[0]], {})
    rows: list[list[str]] = []
    for (r, c), (model, _style) in cells.items():
        if model is None:
            continue
        while len(rows) <= r:
            rows.append([])
        row = rows[r]
        while len(row) <= c:
            row.append("")
        row[c] = csv_text(_to_univer_cell(model))
    return rows
