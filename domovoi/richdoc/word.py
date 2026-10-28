"""A formatted document as Word (.docx), with python-docx (MIT).

Used for "Download as Word", and for saving a ``.docx`` that was opened in
the editor and kept as Word. The input is the editor's content, which the
caller has sanitised (:func:`domovoi.richdoc.sanitize_html`), so the
element set is the editor's own:

* headings, paragraphs and their alignment;
* bold, italic, underline, strike, code, links, text colour, highlight
  (any colour, as cell shading on the run), font family and size,
  superscript and subscript;
* bulleted and numbered lists (three levels of indent, Word's own list
  styles), checklists (as a box character before the item), quotes, code
  blocks, horizontal lines;
* tables (merged cells kept), and pictures with their width, alignment
  and caption.

What Word can't hold the same way (a picture floated beside text is
written on a line of its own, aligned the same way) is said on the save
dialog: "complex formatting may shift".
"""

from __future__ import annotations

import base64
import binascii
import io
import re
from dataclasses import dataclass, replace
from typing import Any

from domovoi.image_safety import LIBRARY_FORMATS, PictureRefused, open_image
from domovoi.richdoc.text import _Node, build_tree

_HEADINGS = {"h1": 1, "h2": 2, "h3": 3, "h4": 4, "h5": 5, "h6": 6}

# What one Word copy may hold. A document is somebody else's input, and the
# conversion runs on the server: each bound keeps the work proportional to
# an ordinary document's (security review round 1: a 60 x 40 table took
# 72 s, and 144 KB of spans reached 1.1 GB). Past them, the rest of that
# table (or the table, or the picture) is left out of the Word copy.
MAX_TABLE_ROWS = 500
MAX_TABLE_COLS = 60
MAX_SPAN = 60
MAX_TABLE_CELLS = 5_000
MAX_DOCUMENT_CELLS = 20_000
MAX_PICTURES = 200
_TEXT_WIDTH_IN = 6.0       # python-docx's default page: 8.5in less 1.25in margins each side
_DATA_URI = re.compile(r"^data:image/([a-z]+);base64,(.*)$", re.I | re.S)
_STYLE_DECL = re.compile(r"\s*([a-z-]+)\s*:\s*([^;]+)")
_NAMED = {
    "black": "000000", "white": "FFFFFF", "red": "FF0000", "green": "008000", "blue": "0000FF",
    "yellow": "FFFF00", "orange": "FFA500", "purple": "800080", "gray": "808080", "grey": "808080",
}


@dataclass(frozen=True)
class _Fmt:
    bold: bool = False
    italic: bool = False
    underline: bool = False
    strike: bool = False
    code: bool = False
    sup: bool = False
    sub: bool = False
    color: str | None = None
    highlight: str | None = None
    font: str | None = None
    size_pt: float | None = None
    href: str | None = None


def _styles(node: _Node) -> dict[str, str]:
    return {m.group(1).lower(): m.group(2).strip() for m in _STYLE_DECL.finditer(node.attrs.get("style", ""))}


def _hex(colour: str | None) -> str | None:
    """CSS colour → RRGGBB, or None for what Word can't take."""
    if not colour:
        return None
    c = colour.strip().lower()
    if c.startswith("#"):
        h = c[1:]
        if len(h) in (3, 4):
            h = "".join(ch * 2 for ch in h[:3])
        return h[:6].upper() if len(h) >= 6 and re.fullmatch(r"[0-9a-f]+", h[:6]) else None
    m = re.match(r"rgba?\(\s*(\d+)\s*,\s*(\d+)\s*,\s*(\d+)", c)
    if m:
        return "".join(f"{min(255, int(x)):02X}" for x in m.groups())
    return _NAMED.get(c)


def _size_pt(value: str | None) -> float | None:
    if not value:
        return None
    m = re.match(r"^([\d.]+)(px|pt|em|rem|%)$", value.strip().lower())
    if not m:
        return None
    n = float(m.group(1))
    unit = m.group(2)
    pt = {"pt": n, "px": n * 0.75, "em": n * 12, "rem": n * 12, "%": n * 0.12}[unit]
    return max(1.0, min(400.0, round(pt * 2) / 2))


def _first_font(family: str | None) -> str | None:
    if not family:
        return None
    first = family.split(",")[0].strip().strip("'\"")
    generic = {"serif": "Times New Roman", "sans-serif": "Arial", "monospace": "Consolas",
               "system-ui": "Calibri", "ui-monospace": "Consolas"}
    return generic.get(first.lower(), first) or None


class _Writer:
    def __init__(self) -> None:
        from docx import Document

        self.doc = Document()
        self.images = 0
        self.cells = 0

    # ── runs ──
    def _add_run(self, paragraph: Any, text: str, fmt: _Fmt) -> None:
        from docx.shared import Pt, RGBColor

        if not text:
            return
        if fmt.href:
            self._add_link(paragraph, text, fmt)
            return
        run = paragraph.add_run(text)
        self._format(run, fmt)

    def _format(self, run: Any, fmt: _Fmt) -> None:
        from docx.oxml import OxmlElement
        from docx.oxml.ns import qn
        from docx.shared import Pt, RGBColor

        run.bold = fmt.bold or None
        run.italic = fmt.italic or None
        run.underline = fmt.underline or None
        if fmt.strike:
            run.font.strike = True
        if fmt.sup:
            run.font.superscript = True
        if fmt.sub:
            run.font.subscript = True
        font = "Consolas" if fmt.code else fmt.font
        if font:
            run.font.name = font
            rpr = run._element.get_or_add_rPr()
            fonts = rpr.find(qn("w:rFonts"))
            if fonts is not None:
                fonts.set(qn("w:eastAsia"), font)
        if fmt.size_pt:
            run.font.size = Pt(fmt.size_pt)
        if fmt.color:
            run.font.color.rgb = RGBColor.from_string(fmt.color)
        if fmt.highlight:
            shd = OxmlElement("w:shd")
            shd.set(qn("w:val"), "clear")
            shd.set(qn("w:color"), "auto")
            shd.set(qn("w:fill"), fmt.highlight)
            run._element.get_or_add_rPr().append(shd)

    def _add_link(self, paragraph: Any, text: str, fmt: _Fmt) -> None:
        from docx.opc.constants import RELATIONSHIP_TYPE as RT
        from docx.oxml import OxmlElement
        from docx.oxml.ns import qn

        href = fmt.href or ""
        if not re.match(r"^(https?:|mailto:|tel:)", href, re.I):
            paragraph.add_run(text)
            return
        r_id = paragraph.part.relate_to(href, RT.HYPERLINK, is_external=True)
        link = OxmlElement("w:hyperlink")
        link.set(qn("r:id"), r_id)
        run = paragraph.add_run(text)
        self._format(run, replace(fmt, href=None, underline=True, color=fmt.color or "1A5FB4"))
        link.append(run._element)
        paragraph._p.append(link)

    def inline(self, paragraph: Any, node: object, fmt: _Fmt) -> None:
        if isinstance(node, str):
            self._add_run(paragraph, node.replace("\r", " ").replace("\n", " "), fmt)
            return
        n: _Node = node  # type: ignore[assignment]
        tag = n.tag
        if tag == "br":
            paragraph.add_run().add_break()
            return
        if tag == "img":
            self.picture_run(paragraph, n, None)
            return
        if tag == "input":
            return
        s = _styles(n)
        if tag in ("strong", "b"):
            fmt = replace(fmt, bold=True)
        elif tag in ("em", "i"):
            fmt = replace(fmt, italic=True)
        elif tag == "u":
            fmt = replace(fmt, underline=True)
        elif tag in ("s", "del", "strike"):
            fmt = replace(fmt, strike=True)
        elif tag == "code":
            fmt = replace(fmt, code=True)
        elif tag == "sup":
            fmt = replace(fmt, sup=True)
        elif tag == "sub":
            fmt = replace(fmt, sub=True)
        elif tag == "a":
            fmt = replace(fmt, href=n.attrs.get("href"))
        elif tag == "mark":
            fmt = replace(fmt, highlight=_hex(s.get("background-color") or n.attrs.get("data-color")) or "FFF3A3")
        if "color" in s:
            fmt = replace(fmt, color=_hex(s["color"]) or fmt.color)
        if "background-color" in s and tag != "mark":
            fmt = replace(fmt, highlight=_hex(s["background-color"]) or fmt.highlight)
        if "font-family" in s:
            fmt = replace(fmt, font=_first_font(s["font-family"]) or fmt.font)
        if "font-size" in s:
            fmt = replace(fmt, size_pt=_size_pt(s["font-size"]) or fmt.size_pt)
        for c in n.children:
            self.inline(paragraph, c, fmt)

    # ── pictures ──
    def _picture_bytes(self, src: str) -> tuple[bytes, int, int] | None:
        m = _DATA_URI.match(src or "")
        if not m:
            return None
        try:
            raw = base64.b64decode(re.sub(r"\s+", "", m.group(2)), validate=True)
        except (binascii.Error, ValueError):
            return None
        try:
            with open_image(io.BytesIO(raw), LIBRARY_FORMATS) as im:
                w, h = im.size
                if im.format in ("PNG", "JPEG", "GIF", "BMP"):
                    return raw, w, h
                out = io.BytesIO()
                im.convert("RGBA" if "A" in im.getbands() else "RGB").save(out, "PNG")
                return out.getvalue(), w, h
        except (PictureRefused, OSError):
            return None

    def picture_run(self, paragraph: Any, img: _Node, width_pct: int | None) -> bool:
        from docx.shared import Inches

        if self.images >= MAX_PICTURES:
            return False
        got = self._picture_bytes(img.attrs.get("src", ""))
        if not got:
            return False
        data, w, h = got
        if width_pct:
            width_in = _TEXT_WIDTH_IN * max(5, min(100, width_pct)) / 100
        else:
            width_in = min(_TEXT_WIDTH_IN, w / 96)
        paragraph.add_run().add_picture(io.BytesIO(data), width=Inches(width_in))
        self.images += 1
        return True

    def figure(self, container: Any, n: _Node) -> None:
        from docx.enum.text import WD_ALIGN_PARAGRAPH

        align = {"left": WD_ALIGN_PARAGRAPH.LEFT, "right": WD_ALIGN_PARAGRAPH.RIGHT}.get(
            n.attrs.get("data-align", ""), WD_ALIGN_PARAGRAPH.CENTER)
        try:
            pct = int(n.attrs.get("data-width") or 0) or None
        except ValueError:
            pct = None
        for c in n.children:
            if isinstance(c, _Node) and c.tag == "img":
                p = container.add_paragraph()
                p.alignment = align
                if not self.picture_run(p, c, pct):
                    self._delete(p)
            elif isinstance(c, _Node) and c.tag == "figcaption":
                p = container.add_paragraph(style="Caption")
                p.alignment = align
                for cc in c.children:
                    self.inline(p, cc, _Fmt())

    @staticmethod
    def _delete(paragraph: Any) -> None:
        el = paragraph._element
        el.getparent().remove(el)

    # ── blocks ──
    def _align(self, paragraph: Any, n: _Node) -> None:
        from docx.enum.text import WD_ALIGN_PARAGRAPH

        a = _styles(n).get("text-align", "").lower()
        mapped = {"center": WD_ALIGN_PARAGRAPH.CENTER, "right": WD_ALIGN_PARAGRAPH.RIGHT,
                  "justify": WD_ALIGN_PARAGRAPH.JUSTIFY, "left": WD_ALIGN_PARAGRAPH.LEFT}.get(a)
        if mapped is not None:
            paragraph.alignment = mapped

    def _para(self, container: Any, style: str | None = None) -> Any:
        try:
            return container.add_paragraph(style=style) if style else container.add_paragraph()
        except KeyError:
            return container.add_paragraph()

    def blocks(self, container: Any, node: _Node, quote: bool = False) -> None:
        run_nodes: list[object] = []

        def flush() -> None:
            if any((isinstance(x, str) and x.strip()) or isinstance(x, _Node) for x in run_nodes):
                p = self._para(container, "Quote" if quote else None)
                for x in run_nodes:
                    self.inline(p, x, _Fmt())
            run_nodes.clear()

        for c in node.children:
            if not isinstance(c, _Node):
                run_nodes.append(c)
                continue
            tag = c.tag
            if tag in _HEADINGS:
                flush()
                p = self.doc.add_heading(level=_HEADINGS[tag]) if container is self.doc \
                    else self._para(container)
                self._align(p, c)
                for cc in c.children:
                    self.inline(p, cc, _Fmt())
            elif tag == "p":
                flush()
                p = self._para(container, "Quote" if quote else None)
                self._align(p, c)
                for cc in c.children:
                    self.inline(p, cc, _Fmt())
            elif tag in ("ul", "ol"):
                flush()
                self.list_(container, c, 1)
            elif tag == "blockquote":
                flush()
                self.blocks(container, c, quote=True)
            elif tag == "pre":
                flush()
                p = self._para(container)
                text = "".join(_raw_text(cc) for cc in c.children).strip("\n")
                for i, line in enumerate(text.split("\n")):
                    if i:
                        p.add_run().add_break()
                    self._add_run(p, line, _Fmt(code=True))
            elif tag == "hr":
                flush()
                self._rule(self._para(container))
            elif tag == "table":
                flush()
                self.table(container, c)
            elif tag == "figure":
                flush()
                self.figure(container, c)
            elif tag in ("div", "section", "article", "main", "figcaption", "li"):
                flush()
                self.blocks(container, c, quote)
            else:
                run_nodes.append(c)
        flush()

    def _rule(self, paragraph: Any) -> None:
        from docx.oxml import OxmlElement
        from docx.oxml.ns import qn

        ppr = paragraph._p.get_or_add_pPr()
        bdr = OxmlElement("w:pBdr")
        bottom = OxmlElement("w:bottom")
        for k, v in (("w:val", "single"), ("w:sz", "6"), ("w:space", "1"), ("w:color", "BFBFBF")):
            bottom.set(qn(k), v)
        bdr.append(bottom)
        ppr.append(bdr)

    def list_(self, container: Any, n: _Node, level: int) -> None:
        numbered = n.tag == "ol"
        task = n.attrs.get("data-type") == "taskList"
        base = "List Number" if numbered and not task else "List Bullet"
        style = base if level == 1 else f"{base} {min(level, 3)}"
        if task:
            style = "List Paragraph" if level == 1 else style
        for li in n.children:
            if not isinstance(li, _Node) or li.tag != "li":
                continue
            p = self._para(container, style)
            if task:
                box = "☑ " if li.attrs.get("data-checked") == "true" else "☐ "
                self._add_run(p, box, _Fmt())
            first = True
            for c in li.children:
                if isinstance(c, _Node) and c.tag in ("ul", "ol"):
                    self.list_(container, c, level + 1)
                elif isinstance(c, _Node) and c.tag in ("p", "div"):
                    if not first:
                        self._add_run(p, " ", _Fmt())
                    self._list_inline(p, c, container, level)
                    first = False
                elif isinstance(c, _Node) and c.tag == "label":
                    continue
                else:
                    self.inline(p, c, _Fmt())

    def _list_inline(self, p: Any, n: _Node, container: Any, level: int) -> None:
        for c in n.children:
            if isinstance(c, _Node) and c.tag in ("ul", "ol"):
                self.list_(container, c, level + 1)
            elif isinstance(c, _Node) and c.tag == "p":
                self._list_inline(p, c, container, level)
            else:
                self.inline(p, c, _Fmt())

    def table(self, container: Any, n: _Node) -> None:
        rows: list[list[_Node]] = []

        def walk(x: _Node) -> None:
            for c in x.children:
                if isinstance(c, _Node) and c.tag == "tr":
                    if len(rows) < MAX_TABLE_ROWS:
                        rows.append([td for td in c.children if isinstance(td, _Node) and td.tag in ("td", "th")])
                elif isinstance(c, _Node) and c.tag in ("thead", "tbody", "tfoot"):
                    walk(c)

        walk(n)
        rows = [r for r in rows if r]
        if not rows:
            return

        def span(td: _Node, key: str) -> int:
            try:
                return max(1, min(MAX_SPAN, int(td.attrs.get(key) or 1)))
            except ValueError:
                return 1

        # Lay the cells out on a grid, as a browser does, so merged cells
        # land where they were. Every span is clamped to the grid BEFORE it
        # is marked taken (columns to MAX_TABLE_COLS, rows to the table), and
        # the grid itself to MAX_TABLE_CELLS, so a hostile colspan/rowspan
        # costs no more than an ordinary table of the same size.
        grid: dict[tuple[int, int], tuple[_Node, int, int]] = {}
        taken: set[tuple[int, int]] = set()
        width = 0
        height = len(rows)
        for r, cells in enumerate(rows):
            col = 0
            for td in cells:
                while (r, col) in taken:
                    col += 1
                if col >= MAX_TABLE_COLS:
                    break
                cs = min(span(td, "colspan"), MAX_TABLE_COLS - col)
                rs = min(span(td, "rowspan"), height - r)
                grid[(r, col)] = (td, cs, rs)
                for dr in range(rs):
                    for dc in range(cs):
                        taken.add((r + dr, col + dc))
                col += cs
                width = max(width, col)
            if len(taken) > MAX_TABLE_CELLS:
                height = r + 1
                break
        if width == 0:
            return
        height = min(height, max(1, MAX_TABLE_CELLS // width))
        if self.cells + width * height > MAX_DOCUMENT_CELLS:
            p = self._para(container)
            self._add_run(p, "[A table too large for a Word copy was left out.]", _Fmt(italic=True))
            return
        self.cells += width * height
        table = container.add_table(rows=height, cols=width)
        try:
            table.style = "Table Grid"
        except (KeyError, ValueError):
            pass
        # Each row's cells fetched ONCE (python-docx's table.cell(r, c)
        # rebuilds the whole cell list on every call: quadratic), the text
        # written first and the merges done after, on those same cells.
        cells = [row.cells for row in table.rows]
        merges: list[tuple[int, int, int, int]] = []
        for (r, c), (td, cs, rs) in grid.items():
            if r >= height or c >= width:
                continue
            cell = cells[r][c]
            end_r, end_c = min(height - 1, r + rs - 1), min(width - 1, c + cs - 1)
            if (end_r, end_c) != (r, c):
                merges.append((r, c, end_r, end_c))
            p = cell.paragraphs[0]
            fmt = _Fmt(bold=td.tag == "th")
            first = True
            for child in td.children:
                if isinstance(child, _Node) and child.tag in ("p", "div"):
                    if not first:
                        p = cell.add_paragraph()
                    self._align(p, child)
                    for cc in child.children:
                        self.inline(p, cc, fmt)
                    first = False
                elif isinstance(child, _Node) and child.tag in ("ul", "ol"):
                    self.list_(cell, child, 1)
                else:
                    self.inline(p, child, fmt)
        for r, c, end_r, end_c in merges:
            try:
                cells[r][c].merge(cells[end_r][end_c])
            except Exception:  # noqa: BLE001 - overlapping spans: keep the cells unmerged
                continue


def _raw_text(node: object) -> str:
    if isinstance(node, str):
        return node
    n: _Node = node  # type: ignore[assignment]
    if n.tag == "br":
        return "\n"
    return "".join(_raw_text(c) for c in n.children)


def html_to_docx(body_html: str, title: str = "") -> bytes:
    """The sanitised editor content ``body_html`` as .docx bytes."""
    w = _Writer()
    if title:
        w.doc.core_properties.title = title[:200]
    w.blocks(w.doc, build_tree(body_html or ""))
    buf = io.BytesIO()
    w.doc.save(buf)
    return buf.getvalue()
