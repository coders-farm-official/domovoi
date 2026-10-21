"""A draft as a .docx: python-docx, the template's sizes and spacing, and
a font that is *named*, never embedded (Arial for the sans templates,
Times New Roman for the serif ones: Liberation Sans and Serif are their
metric twins, so the .docx and the PDF break lines alike). Word or the
reader's own font draws it.

Standard structure only, which is what applicant tracking systems read
best: the name as a large bold line in the body (not a page header), each
section a real "Heading 2", bullets in Word's own "List Bullet" style, an
address block as one paragraph with line breaks. No tables, text boxes or
columns. The phone's twin is ``DraftDocx.kt``.
"""

from __future__ import annotations

import io
import unicodedata

from domovoi.drafts import blocks as blocks_mod
from domovoi.drafts.templates import Style, Template


def clean(text: str) -> str:
    return " ".join(unicodedata.normalize("NFKC", text or "").split())


def _font(run_or_style, name: str, size: float, bold: bool) -> None:
    from docx.oxml.ns import qn
    from docx.shared import Pt, RGBColor

    f = run_or_style.font
    f.name = name
    f.size = Pt(size)
    f.bold = bold
    f.color.rgb = RGBColor(0, 0, 0)
    rpr = run_or_style.element.get_or_add_rPr()
    fonts = rpr.find(qn("w:rFonts"))
    if fonts is None:
        fonts = rpr.makeelement(qn("w:rFonts"), {})
        rpr.append(fonts)
    for attr in ("w:ascii", "w:hAnsi", "w:cs", "w:eastAsia"):
        fonts.set(qn(attr), name)


def _spacing(paragraph, st: Style) -> None:
    from docx.shared import Pt

    pf = paragraph.paragraph_format
    pf.space_before = Pt(st.before)
    pf.space_after = Pt(st.after)
    pf.line_spacing = Pt(st.leading)


def render(markdown: str, template: Template, *, title: str = "") -> bytes:
    from docx import Document
    from docx.shared import Pt

    doc = Document()
    for section in doc.sections:
        section.page_width = Pt(template.width)
        section.page_height = Pt(template.height)
        m = Pt(template.margin)
        section.left_margin = section.right_margin = section.top_margin = section.bottom_margin = m
    body = template.styles["body"]
    _font(doc.styles["Normal"], template.docx_font, body.size, False)
    heading = template.styles["heading"]
    _font(doc.styles["Heading 2"], template.docx_font, heading.size, True)
    _font(doc.styles["List Bullet"], template.docx_font, template.styles["bullet"].size, False)
    doc.core_properties.title = title or ""
    doc.core_properties.author = ""
    doc.core_properties.comments = ""

    for block in blocks_mod.parse(markdown):
        st = template.styles[block.style]
        lines = [clean(x) for x in block.lines if clean(x)]
        if not lines:
            continue
        if block.style == "heading":
            p = doc.add_paragraph(style="Heading 2")
        elif block.style == "bullet":
            p = doc.add_paragraph(style="List Bullet")
        else:
            p = doc.add_paragraph()
        _spacing(p, st)
        bold = st.font == "bold"
        for i, line in enumerate(lines):
            run = p.add_run(line)
            _font(run, template.docx_font, st.size, bold)
            if i < len(lines) - 1:
                run.add_break()
    buf = io.BytesIO()
    doc.save(buf)
    return buf.getvalue()
