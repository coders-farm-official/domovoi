"""The Univer sheet editor's workbook model (web/backend/api/sheet_workbook.py)
and its routes (GET/PUT /api/documents/workbook, POST /api/documents/presence).

Univer's own .xlsx import/export is in its paid tier, so the server
converts with openpyxl. These pin the fidelity of that round trip on sample
workbooks made here: what the editor shows comes back exactly when nothing
was edited, an edit changes only what was edited, and what the editor
doesn't model (comments, data validation, conditional formatting) is kept.
"""

from __future__ import annotations

import codecs
import copy
import datetime as dt
import json

import pytest
from fastapi.testclient import TestClient
from openpyxl import Workbook, load_workbook
from openpyxl.chart import BarChart, Reference
from openpyxl.comments import Comment
from openpyxl.formatting.rule import CellIsRule
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.worksheet.datavalidation import DataValidation

from domovoi.config import settings
from domovoi.tests.auth_testkit import install_fake_db
from web.backend.api import documents as docs
from web.backend.api import sheet_workbook as sw
from web.backend.main import app


@pytest.fixture(autouse=True)
def _pre_setup_install(monkeypatch):
    # Document behaviour, not auth (as test_documents_sheet_preserve.py).
    install_fake_db(monkeypatch, admin=False)


@pytest.fixture
def docs_dir(monkeypatch, tmp_path):
    monkeypatch.setattr(settings, "documents_dir", str(tmp_path))
    docs._presence.clear()
    return tmp_path


def _client():
    return TestClient(app, headers={"X-Requested-With": "domovoi-tests"})


# ─── sample workbooks ───────────────────────────────────────────────
def _household(path):
    """Values, formulas, dates, number formats, fonts, fills, borders,
    alignment, merges, sizes, freeze, three sheets (one hidden), and the
    things the editor doesn't model: a comment, data validation and
    conditional formatting."""
    wb = Workbook()
    ws = wb.active
    ws.title = "Budget"
    ws["A1"] = "Item"
    ws["B1"] = "Cost"
    ws["C1"] = "Paid on"
    for ref in ("A1", "B1", "C1"):
        ws[ref].font = Font(bold=True, color="FFFFFFFF", name="Arial", size=12)
        ws[ref].fill = PatternFill("solid", fgColor="FF336699")
        ws[ref].alignment = Alignment(horizontal="center", vertical="center")
    ws["A2"], ws["B2"], ws["C2"] = "Rent", 1200, dt.datetime(2026, 10, 1)
    ws["A3"], ws["B3"], ws["C3"] = "Food", 412.5, dt.date(2026, 10, 3)
    ws["B2"].number_format = ws["B3"].number_format = '"$"#,##0.00'
    ws["C2"].number_format = ws["C3"].number_format = "yyyy-mm-dd"
    ws["A4"] = "Total"
    ws["B4"] = "=SUM(B2:B3)"
    ws["A4"].font = Font(italic=True, underline="single", strike=True)
    thin = Side(style="thin", color="FF000000")
    ws["B4"].border = Border(top=thin, bottom=Side(style="double", color="FFFF0000"))
    ws["A6"] = "A long note that wraps inside its merged cells"
    ws["A6"].alignment = Alignment(wrap_text=True, vertical="top")
    ws.merge_cells("A6:C7")
    ws["D1"] = True
    ws["D2"] = "007"
    ws["E1"] = "rotated"
    ws["E1"].alignment = Alignment(text_rotation=45)
    ws.column_dimensions["A"].width = 24
    ws.column_dimensions["C"].width = 14
    ws.row_dimensions[1].height = 30
    ws.freeze_panes = "A2"
    ws["A2"].comment = Comment("monthly", "Kamron")
    dv = DataValidation(type="list", formula1='"yes,no"')
    ws.add_data_validation(dv)
    dv.add("D5")
    ws.conditional_formatting.add(
        "B2:B3", CellIsRule(operator="greaterThan", formula=["1000"], fill=PatternFill("solid", fgColor="FFFFC7CE"))
    )
    notes = wb.create_sheet("Notes")
    notes["A1"] = "keep me"
    notes.sheet_properties.tabColor = "FF00AA00"
    secret = wb.create_sheet("Hidden")
    secret["A1"] = 42
    secret.sheet_state = "hidden"
    wb.save(path)


def _cell_facts(c):
    f, fl, al, bd = c.font, c.fill, c.alignment, c.border
    return {
        "value": c.value,
        "fmt": c.number_format,
        "font": (f.name, f.sz, f.b, f.i, f.u, f.strike, sw._openpyxl_hex(f.color)),
        "fill": (fl.fill_type, sw._openpyxl_hex(fl.fgColor) if fl.fill_type else None),
        "align": (al.horizontal, al.vertical, al.wrap_text, al.text_rotation),
        "border": tuple((s.style, sw._openpyxl_hex(s.color) if s.style else None) if s else None
                        for s in (bd.top, bd.bottom, bd.left, bd.right)),
    }


def _facts(path):
    wb = load_workbook(path)
    out = {"sheets": [(ws.title, ws.sheet_state) for ws in wb.worksheets]}
    for ws in wb.worksheets:
        cells = {}
        for row in ws.iter_rows():
            for c in row:
                if c.value is not None or c.has_style:
                    cells[c.coordinate] = _cell_facts(c)
        out[ws.title] = {
            "cells": cells,
            "merges": sorted(str(r) for r in ws.merged_cells.ranges),
            "widths": {k: d.width for k, d in ws.column_dimensions.items() if d.width},
            "heights": {k: d.height for k, d in ws.row_dimensions.items() if d.height},
            "freeze": ws.freeze_panes,
            "tab": sw._openpyxl_hex(ws.sheet_properties.tabColor) if ws.sheet_properties.tabColor else None,
            "comments": {c.coordinate: c.comment.text for row in ws.iter_rows() for c in row if c.comment},
            "validations": [(d.type, d.formula1, str(d.sqref)) for d in ws.data_validations.dataValidation],
            "cf": sorted(str(rng.sqref) for rng in ws.conditional_formatting),
        }
    return out


def _snapshot(path):
    """What the editor gets, then sends back untouched (JSON round trip,
    as the browser does)."""
    model = sw.read_xlsx(path)
    wb = json.loads(json.dumps(model["workbook"]))
    for sid, sheet in wb["sheets"].items():
        sheet["origin"] = sid
    return wb


def _save(path, snapshot):
    wb = load_workbook(path)
    sw.write_xlsx(wb, snapshot)
    wb.save(path)


# ─── fidelity: read → unchanged save is a no-op ─────────────────────
def test_reading_shows_what_the_file_holds(tmp_path):
    _household(tmp_path / "h.xlsx")
    model = sw.read_xlsx(tmp_path / "h.xlsx")
    wb = json.loads(json.dumps(model["workbook"]))  # as the browser gets it
    assert [wb["sheets"][s]["name"] for s in wb["sheetOrder"]] == ["Budget", "Notes", "Hidden"]
    budget, notes, hidden = (wb["sheets"][s] for s in wb["sheetOrder"])
    cells = budget["cellData"]
    style = lambda cell: wb["styles"][cell["s"]]  # noqa: E731
    assert cells["0"]["0"]["v"] == "Item" and style(cells["0"]["0"])["bl"] == 1
    assert style(cells["0"]["0"])["bg"] == {"rgb": "#336699"}
    assert style(cells["0"]["0"])["cl"] == {"rgb": "#FFFFFF"}
    assert style(cells["0"]["0"])["ff"] == "Arial" and style(cells["0"]["0"])["fs"] == 12
    assert style(cells["0"]["0"])["ht"] == 2 and style(cells["0"]["0"])["vt"] == 2
    assert cells["1"]["1"] == {"v": 1200, "t": 2, "s": cells["1"]["1"]["s"]}
    assert style(cells["1"]["1"])["n"] == {"pattern": '"$"#,##0.00'}
    # Dates are serial numbers with their date format (Univer shows them as dates).
    assert cells["1"]["2"]["v"] == 46296 and style(cells["1"]["2"])["n"] == {"pattern": "yyyy-mm-dd"}
    assert cells["3"]["1"]["f"] == "=SUM(B2:B3)"
    assert style(cells["3"]["0"])["it"] == 1 and style(cells["3"]["0"])["ul"] == {"s": 1}
    assert style(cells["3"]["0"])["st"] == {"s": 1}
    assert style(cells["3"]["1"])["bd"]["b"] == {"s": 7, "cl": {"rgb": "#FF0000"}}
    assert cells["0"]["3"] == {"v": 1, "t": 3}
    assert cells["1"]["3"] == {"v": "007", "t": 1}
    assert style(cells["0"]["4"])["tr"] == {"a": -45}   # Excel 45 = up = Univer -45
    assert budget["mergeData"] == [{"startRow": 5, "endRow": 6, "startColumn": 0, "endColumn": 2}]
    assert budget["columnData"]["0"]["w"] == sw.width_to_px(24)
    assert budget["rowData"]["0"]["h"] == 40
    assert budget["freeze"] == {"xSplit": 0, "ySplit": 1, "startRow": 1, "startColumn": 0}
    assert notes["tabColor"] == "#00AA00"
    assert hidden["hidden"] == 1 and hidden["cellData"]["0"]["0"] == {"v": 42, "t": 2}
    assert model["truncated"] == []


def test_an_unchanged_save_changes_nothing(tmp_path):
    path = tmp_path / "h.xlsx"
    _household(path)
    before = _facts(path)
    _save(path, _snapshot(path))
    assert _facts(path) == before


# ─── edits change only what was edited ──────────────────────────────
def test_an_edit_changes_that_cell_and_keeps_everything_else(tmp_path):
    path = tmp_path / "h.xlsx"
    _household(path)
    before = _facts(path)
    snap = _snapshot(path)
    budget = snap["sheets"]["s0"]
    budget["cellData"]["2"]["1"]["v"] = 450           # Food costs more
    budget["cellData"].setdefault("8", {})["0"] = {"v": "new row", "t": 1}
    _save(path, snap)
    after = _facts(path)
    assert after["Budget"]["cells"]["B3"]["value"] == 450
    assert after["Budget"]["cells"]["B3"]["fmt"] == '"$"#,##0.00'   # format kept
    assert after["Budget"]["cells"]["A9"]["value"] == "new row"
    del after["Budget"]["cells"]["A9"]
    after["Budget"]["cells"]["B3"]["value"] = 412.5
    assert after == before


def test_formatting_edits_are_written(tmp_path):
    path = tmp_path / "h.xlsx"
    _household(path)
    snap = _snapshot(path)
    sheet = snap["sheets"]["s0"]
    snap["styles"]["new1"] = {
        "bl": 1, "ff": "Inter", "fs": 14, "cl": {"rgb": "rgb(255, 0, 0)"},
        "bg": {"rgb": "#ffff00"}, "ht": 3, "tb": 3,
        "bd": {"l": {"s": 8, "cl": {"rgb": "#0000FF"}}},
        "n": {"pattern": "0.0%"},
    }
    sheet["cellData"]["2"]["0"]["s"] = "new1"        # A3 "Food"
    sheet["cellData"]["1"]["0"]["s"] = None          # A2 "Rent" loses its style (it had none)
    _save(path, snap)
    c = load_workbook(path)["Budget"]["A3"]
    assert c.value == "Food"
    assert (c.font.b, c.font.name, c.font.sz) == (True, "Inter", 14)
    assert sw._openpyxl_hex(c.font.color) == "#FF0000"
    assert c.fill.fill_type == "solid" and sw._openpyxl_hex(c.fill.fgColor) == "#FFFF00"
    assert c.alignment.horizontal == "right" and c.alignment.wrap_text
    assert c.border.left.style == "medium" and sw._openpyxl_hex(c.border.left.color) == "#0000FF"
    assert c.number_format == "0.0%"


def test_clearing_a_cell_and_unbolding_a_header(tmp_path):
    path = tmp_path / "h.xlsx"
    _household(path)
    snap = _snapshot(path)
    sheet = snap["sheets"]["s0"]
    del sheet["cellData"]["2"]["0"]                     # A3 cleared
    hdr = dict(snap["styles"][sheet["cellData"]["0"]["1"]["s"]])
    hdr.pop("bl")
    snap["styles"]["plain"] = hdr
    sheet["cellData"]["0"]["1"]["s"] = "plain"         # B1 no longer bold
    _save(path, snap)
    ws = load_workbook(path)["Budget"]
    assert ws["A3"].value is None
    assert ws["B1"].font.b is False and ws["B1"].value == "Cost"
    assert sw._openpyxl_hex(ws["B1"].fill.fgColor) == "#336699"   # fill kept


def test_merges_sizes_and_freeze(tmp_path):
    path = tmp_path / "h.xlsx"
    _household(path)
    snap = _snapshot(path)
    sheet = snap["sheets"]["s0"]
    sheet["mergeData"] = [{"startRow": 9, "endRow": 9, "startColumn": 0, "endColumn": 1}]
    sheet["columnData"] = {0: {"w": 200}, 3: {"w": 40, "hd": 1}}   # C loses its width
    sheet["rowData"] = {0: {"h": 40}, 4: {"h": 60}}
    sheet["freeze"] = {"xSplit": 1, "ySplit": 2, "startRow": 2, "startColumn": 1}
    _save(path, snap)
    ws = load_workbook(path)["Budget"]
    assert sorted(str(r) for r in ws.merged_cells.ranges) == ["A10:B10"]
    assert ws["A6"].value.startswith("A long note")            # unmerged, value kept
    assert ws.column_dimensions["A"].width == sw.px_to_width(200)
    assert ws.column_dimensions["D"].hidden and ws.column_dimensions["D"].width == sw.px_to_width(40)
    assert "C" not in ws.column_dimensions or not ws.column_dimensions["C"].width
    assert ws.row_dimensions[5].height == 45
    assert ws.freeze_panes == "B3"


def test_sheets_added_renamed_reordered_hidden_and_deleted(tmp_path):
    path = tmp_path / "h.xlsx"
    _household(path)
    snap = _snapshot(path)
    snap["sheets"]["s1"]["name"] = "Budget"          # swap names with s0
    snap["sheets"]["s0"]["name"] = "Notes"
    snap["sheets"]["s2"]["hidden"] = 0
    snap["sheets"]["newX"] = {"id": "newX", "name": "Plans", "origin": None,
                              "cellData": {0: {0: {"v": "trip", "t": 1}}}}
    snap["sheetOrder"] = ["newX", "s2", "s0"]       # Notes (s1) deleted
    _save(path, snap)
    wb = load_workbook(path)
    assert [(ws.title, ws.sheet_state) for ws in wb.worksheets] == [
        ("Plans", "visible"), ("Hidden", "visible"), ("Notes", "visible"),
    ]
    assert wb["Plans"]["A1"].value == "trip"
    assert wb["Notes"]["A1"].value == "Item"             # the old Budget, renamed
    assert wb["Notes"]["A2"].comment.text == "monthly"
    # The next read gives the new order's ids, which the next save uses.
    again = sw.read_xlsx(path)["workbook"]
    assert [again["sheets"][s]["name"] for s in again["sheetOrder"]] == ["Plans", "Hidden", "Notes"]


def test_a_date_cell_edited_stays_a_date(tmp_path):
    path = tmp_path / "h.xlsx"
    _household(path)
    snap = _snapshot(path)
    snap["sheets"]["s0"]["cellData"]["1"]["2"]["v"] = 46300   # 2026-10-05
    _save(path, snap)
    c = load_workbook(path)["Budget"]["C2"]
    assert c.value == dt.datetime(2026, 10, 5) and c.number_format == "yyyy-mm-dd"


def test_a_formula_typed_and_text_forced(tmp_path):
    path = tmp_path / "h.xlsx"
    _household(path)
    snap = _snapshot(path)
    cd = snap["sheets"]["s0"]["cellData"]
    cd["3"]["2"] = {"f": "=B4*2", "v": 3225}          # Univer sends the result too
    cd.setdefault("10", {})["0"] = {"v": "=not a formula", "t": 4}
    cd["10"]["1"] = {"v": "TRUE", "t": 3}
    _save(path, snap)
    ws = load_workbook(path)["Budget"]
    assert ws["C4"].value == "=B4*2"
    assert ws["A11"].value == "=not a formula" and ws["A11"].data_type == "s"
    assert ws["B11"].value is True


def test_text_turned_up_or_down_keeps_its_direction(tmp_path):
    # Univer turns text up with a negative angle, Excel with 1-90.
    path = tmp_path / "h.xlsx"
    _household(path)
    snap = _snapshot(path)
    snap["styles"]["up30"] = {"tr": {"a": -30}}
    snap["styles"]["down20"] = {"tr": {"a": 20}}
    snap["styles"]["stacked"] = {"tr": {"a": 0, "v": 1}}
    cd = snap["sheets"]["s0"]["cellData"]
    cd.setdefault("11", {})["0"] = {"v": "up", "t": 1, "s": "up30"}
    cd["11"]["1"] = {"v": "down", "t": 1, "s": "down20"}
    cd["11"]["2"] = {"v": "stacked", "t": 1, "s": "stacked"}
    _save(path, snap)
    ws = load_workbook(path)["Budget"]
    assert ws["A12"].alignment.text_rotation == 30
    assert ws["B12"].alignment.text_rotation == 110
    assert ws["C12"].alignment.text_rotation == 255
    again = json.loads(json.dumps(sw.read_xlsx(path)["workbook"]))
    styles, row = again["styles"], again["sheets"]["s0"]["cellData"]["11"]
    assert [styles[row[c]["s"]]["tr"] for c in ("0", "1", "2")] == [{"a": -30}, {"a": 20}, {"a": 0, "v": 1}]


def test_style_spellings_univer_may_use_compare_equal():
    a = sw.norm_style({"bl": 1, "cl": {"rgb": "rgb(0,0,0)"}, "bg": {"rgb": "#abc"}, "ht": 0, "ul": {"s": 0}})
    assert a == {"bl": 1, "bg": {"rgb": "#AABBCC"}}
    assert sw.norm_style({"tr": {"a": 45, "v": 0}}) == {"tr": {"a": 45}}
    assert sw.norm_style(None) == {}


def test_a_cell_past_the_window_is_refused_not_dropped(tmp_path):
    path = tmp_path / "h.xlsx"
    _household(path)
    snap = _snapshot(path)
    snap["sheets"]["s0"]["cellData"][sw.WINDOW_ROWS] = {0: {"v": "too far", "t": 1}}
    with pytest.raises(sw.WorkbookRefused) as err:
        _save(path, snap)
    assert err.value.status_code == 413


def test_a_sheet_larger_than_the_window_keeps_what_is_past_it(tmp_path):
    path = tmp_path / "big.xlsx"
    wb = Workbook()
    ws = wb.active
    ws["A1"] = "top"
    ws.cell(row=sw.WINDOW_ROWS + 5, column=1, value="past the window")
    wb.save(path)
    model = sw.read_xlsx(path)
    assert model["truncated"] == ["Sheet"]
    snap = json.loads(json.dumps(model["workbook"]))
    snap["sheets"]["s0"]["origin"] = "s0"
    snap["sheets"]["s0"]["cellData"]["0"]["0"]["v"] = "changed"
    _save(path, snap)
    ws = load_workbook(path).active
    assert ws["A1"].value == "changed"
    assert ws.cell(row=sw.WINDOW_ROWS + 5, column=1).value == "past the window"


def test_grouped_column_widths_are_split_before_one_changes(tmp_path):
    path = tmp_path / "g.xlsx"
    wb = Workbook()
    ws = wb.active
    ws["A1"] = "x"
    ws.column_dimensions.group("B", "D", hidden=False)
    ws.column_dimensions["B"].width = 20
    wb.save(path)
    snap = _snapshot(path)
    cols = snap["sheets"]["s0"]["columnData"]
    assert {int(k) for k in cols} == {1, 2, 3}
    cols["2"] = {"w": 300}
    _save(path, snap)
    dims = load_workbook(path).active.column_dimensions
    spans = sorted((d.min, d.max) for d in dims.values() if d.width)
    assert spans == [(2, 2), (3, 3), (4, 4)]               # no overlapping <col>
    assert dims["C"].width == sw.px_to_width(300)
    assert dims["B"].width == dims["D"].width == 20


# ─── .csv ───────────────────────────────────────────────────────────
def test_csv_numbers_only_when_they_read_back_the_same(tmp_path):
    path = tmp_path / "c.csv"
    path.write_bytes(codecs.BOM_UTF8 + b"name,qty,code,price\nbolt,1200,007,1.50\nnut,3,=B2*2,2.5\n")
    wb = json.loads(json.dumps(sw.read_csv(path, 100, 10)["workbook"]))
    cd = wb["sheets"]["s0"]["cellData"]
    assert cd["1"]["1"] == {"v": 1200, "t": 2}
    assert cd["1"]["2"] == {"v": "007", "t": 1}
    assert cd["1"]["3"] == {"v": "1.50", "t": 1}
    assert cd["2"]["3"] == {"v": 2.5, "t": 2}
    assert cd["2"]["2"] == {"f": "=B2*2"}
    rows = sw.csv_rows(wb)
    assert rows == [["name", "qty", "code", "price"], ["bolt", "1200", "007", "1.50"],
                    ["nut", "3", "=B2*2", "2.5"]]


# ─── the routes ─────────────────────────────────────────────────────
def test_get_and_put_round_trip_with_the_version_check(docs_dir):
    _household(docs_dir / "h.xlsx")
    before = _facts(docs_dir / "h.xlsx")
    with _client() as c:
        r = c.get("/api/documents/workbook/h.xlsx")
        assert r.status_code == 200, r.text
        body = r.json()
        assert body["editable"] is True and body["format"] == "xlsx" and body["version"]
        snap = body["workbook"]
        for sid, sheet in snap["sheets"].items():
            sheet["origin"] = sid
        r2 = c.put("/api/documents/workbook/h.xlsx", json={"workbook": snap, "base_version": body["version"]})
        assert r2.status_code == 200, r2.text
        assert r2.json()["version"] != body["version"]
        assert _facts(docs_dir / "h.xlsx") == before
        # A second save from the stale version is someone else's overwrite.
        stale = c.put("/api/documents/workbook/h.xlsx", json={"workbook": snap, "base_version": body["version"]})
        assert stale.status_code == 409
        assert stale.json()["detail"]["reason"] == "changed"
        forced = c.put("/api/documents/workbook/h.xlsx",
                       json={"workbook": snap, "base_version": body["version"], "force": True})
        assert forced.status_code == 200


def test_a_file_with_a_chart_opens_read_only_and_is_not_saved(docs_dir):
    path = docs_dir / "chart.xlsx"
    wb = Workbook()
    ws = wb.active
    for i in range(1, 4):
        ws.cell(row=i, column=1, value=i)
    chart = BarChart()
    chart.add_data(Reference(ws, min_col=1, min_row=1, max_row=3))
    ws.add_chart(chart, "C1")
    wb.save(path)
    raw = path.read_bytes()
    with _client() as c:
        body = c.get("/api/documents/workbook/chart.xlsx").json()
        assert body["editable"] is False and "charts" in body["why"]
        snap = body["workbook"]
        r = c.put("/api/documents/workbook/chart.xlsx", json={"workbook": snap, "base_version": body["version"]})
    assert r.status_code == 409
    assert path.read_bytes() == raw


def test_csv_route_keeps_its_bom(docs_dir):
    path = docs_dir / "list.csv"
    path.write_bytes(codecs.BOM_UTF8 + b"a,b\n1,2\n")
    with _client() as c:
        body = c.get("/api/documents/workbook/list.csv").json()
        assert body["format"] == "csv"
        snap = body["workbook"]
        snap["sheets"]["s0"]["cellData"]["1"]["1"] = {"v": 5, "t": 2}
        r = c.put("/api/documents/workbook/list.csv", json={"workbook": snap, "base_version": body["version"]})
        assert r.status_code == 200, r.text
    assert path.read_bytes() == codecs.BOM_UTF8 + b"a,b\r\n1,5\r\n"


def test_an_empty_new_xlsx_opens_blank_and_saves(docs_dir):
    (docs_dir / "new.xlsx").write_bytes(b"")
    with _client() as c:
        body = c.get("/api/documents/workbook/new.xlsx").json()
        snap = body["workbook"]
        snap["sheets"]["s0"]["cellData"] = {"0": {"0": {"v": "hello", "t": 1}}}
        r = c.put("/api/documents/workbook/new.xlsx", json={"workbook": snap, "base_version": body["version"]})
        assert r.status_code == 200, r.text
    ws = load_workbook(docs_dir / "new.xlsx").active
    assert ws["A1"].value == "hello" and ws.title == "Sheet1"


def test_other_types_are_refused(docs_dir):
    (docs_dir / "notes.md").write_text("# hi", encoding="utf-8")
    (docs_dir / "old.ods").write_bytes(b"PK")
    with _client() as c:
        assert c.get("/api/documents/workbook/old.ods").status_code == 415
        r = c.put("/api/documents/workbook/notes.md",
                  json={"workbook": {"sheetOrder": [], "sheets": {}}, "base_version": None})
        assert r.status_code == 415
    assert (docs_dir / "notes.md").read_text(encoding="utf-8") == "# hi"


def test_presence_names_the_other_editors(docs_dir):
    _household(docs_dir / "h.xlsx")
    with _client() as c:
        a = c.post("/api/documents/presence/h.xlsx", json={"session_id": "session-aaaa", "device_id": "dev-a"})
        assert a.status_code == 200, a.text
        assert a.json()["others"] == []
        b = c.post("/api/documents/presence/h.xlsx", json={"session_id": "session-bbbb", "device_id": "dev-b"})
        assert b.json()["others"] == ["another device"]   # no devices row in the fake DB
        c.post("/api/documents/presence/h.xlsx", json={"session_id": "session-bbbb", "leaving": True})
        a2 = c.post("/api/documents/presence/h.xlsx", json={"session_id": "session-aaaa", "device_id": "dev-a"})
        assert a2.json()["others"] == []


# ─── the server's other readers still read what the editor saved ────
def test_chat_extract_reads_a_workbook_the_editor_saved(tmp_path):
    from domovoi.doc_extract import sheets as extract

    path = tmp_path / "h.xlsx"
    _household(path)
    snap = _snapshot(path)
    snap["sheets"]["s0"]["cellData"]["2"]["1"]["v"] = 450
    _save(path, snap)
    fn = getattr(extract, "extract_xlsx", None) or getattr(extract, "extract")
    text = json.dumps(fn(path), default=str) if not isinstance(fn(path), str) else fn(path)
    assert "Rent" in text and "450" in text and "keep me" in text
