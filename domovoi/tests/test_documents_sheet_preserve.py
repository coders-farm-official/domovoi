"""The sheet editor's save keeps what it cannot show (web/backend/api/documents.py).

The save used to build a brand-new workbook from the editor's grid, so one
save in the dashboard (or the app's connected sheet editor) of a household
.xlsx threw away its formatting, column widths, merges, every sheet but the
first and anything past the editor's 1000 x 60 window. These pin the
in-place save that replaced it, the 409 for parts openpyxl cannot carry,
the CSV equivalent, and the read no longer padding every row to 60 columns.
"""

from __future__ import annotations

import codecs
import datetime as dt

import pytest
from fastapi.testclient import TestClient
from openpyxl import Workbook, load_workbook
from openpyxl.chart import BarChart, Reference
from openpyxl.styles import Font, PatternFill

from domovoi.config import settings
from domovoi.tests.auth_testkit import install_fake_db
from web.backend.api import documents as docs
from web.backend.main import app


@pytest.fixture(autouse=True)
def _pre_setup_install(monkeypatch):
    # Document behaviour, not auth: a fresh install's pre-setup grace (see
    # test_documents.py); who may call what is proven in test_web_media_auth.py.
    install_fake_db(monkeypatch, admin=False)


@pytest.fixture
def docs_dir(monkeypatch, tmp_path):
    monkeypatch.setattr(settings, "documents_dir", str(tmp_path))
    return tmp_path


def _client():
    return TestClient(app, headers={"X-Requested-With": "domovoi-tests"})


def _budget(path):
    """A household workbook with everything the editor cannot show."""
    wb = Workbook()
    ws = wb.active
    ws.title = "Budget"
    ws["A1"] = "Item"
    ws["B1"] = "Cost"
    ws["A1"].font = Font(bold=True, color="FFFF0000")
    ws["B1"].fill = PatternFill("solid", fgColor="FFFFFF00")
    ws["A2"] = "Bread"
    ws["B2"] = 3.5
    ws["B2"].number_format = "0.00"
    ws["A3"] = "Paid on"
    ws["B3"] = dt.datetime(2026, 10, 1)
    ws["A4"] = "Done"
    ws["B4"] = True
    ws["C2"] = "=B2*2"
    ws.column_dimensions["A"].width = 30
    ws.merge_cells("A6:C6")
    ws["A6"] = "merged note"
    ws.cell(row=1001, column=1, value="past the last row")
    ws.cell(row=2, column=61, value="past the last column")
    other = wb.create_sheet("Notes")
    other["A1"] = "keep me"
    wb.save(path)


def _grid(c, name):
    r = c.get(f"/api/documents/sheet/{name}")
    assert r.status_code == 200, r.text
    return r.json()["rows"]


def _put(c, name, rows):
    return c.put(f"/api/documents/sheet/{name}", json={"rows": rows})


# ─── .xlsx: an edit changes the edited cell and nothing else ────────
def test_a_save_keeps_formatting_widths_merges_and_other_sheets(docs_dir):
    _budget(docs_dir / "budget.xlsx")
    with _client() as c:
        rows = _grid(c, "budget.xlsx")
        rows[1][1] = {"v": "4.25"}            # Bread costs more
        assert _put(c, "budget.xlsx", rows).status_code == 200

    wb = load_workbook(docs_dir / "budget.xlsx")
    ws = wb["Budget"]
    assert ws["B2"].value == 4.25
    assert ws["B2"].number_format == "0.00"
    assert ws["A1"].font.bold and ws["A1"].font.color.rgb == "FFFF0000"
    assert ws["B1"].fill.fgColor.rgb == "FFFFFF00"
    assert ws.column_dimensions["A"].width == 30
    assert "A6:C6" in {str(m) for m in ws.merged_cells.ranges}
    assert ws["A6"].value == "merged note"
    assert wb.sheetnames == ["Budget", "Notes"]
    assert wb["Notes"]["A1"].value == "keep me"


def test_untouched_dates_booleans_and_formulas_keep_their_types(docs_dir):
    # The editor gets str(value) for a date and a boolean; sending that back
    # unchanged must not turn them into text.
    _budget(docs_dir / "budget.xlsx")
    with _client() as c:
        rows = _grid(c, "budget.xlsx")
        assert rows[2][1]["v"] == "2026-10-01 00:00:00"
        assert rows[3][1]["v"] == "True"
        assert rows[1][2]["f"] == "=B2*2"
        rows[0][0] = {"v": "Thing"}
        assert _put(c, "budget.xlsx", rows).status_code == 200

    ws = load_workbook(docs_dir / "budget.xlsx")["Budget"]
    assert ws["B3"].value == dt.datetime(2026, 10, 1)
    assert ws["B4"].value is True
    assert ws["C2"].value == "=B2*2"
    assert ws["A1"].value == "Thing"


def test_cells_past_the_editors_window_survive(docs_dir):
    _budget(docs_dir / "budget.xlsx")
    with _client() as c:
        rows = _grid(c, "budget.xlsx")
        assert len(rows) <= docs._SHEET_MAX_ROWS
        assert all(len(r) <= docs._SHEET_MAX_COLS for r in rows)
        assert _put(c, "budget.xlsx", rows).status_code == 200

    ws = load_workbook(docs_dir / "budget.xlsx")["Budget"]
    assert ws.cell(row=1001, column=1).value == "past the last row"
    assert ws.cell(row=2, column=61).value == "past the last column"


def test_a_cleared_cell_loses_its_value_but_keeps_its_formatting(docs_dir):
    _budget(docs_dir / "budget.xlsx")
    with _client() as c:
        rows = _grid(c, "budget.xlsx")
        rows[0][0] = None                      # cleared in the editor
        rows[1] = rows[1][:1]                  # trailing cells dropped by the client
        assert _put(c, "budget.xlsx", rows).status_code == 200

    ws = load_workbook(docs_dir / "budget.xlsx")["Budget"]
    assert ws["A1"].value is None
    assert ws["A1"].font.bold                  # the style stayed on the empty cell
    assert ws["B2"].value is None and ws["C2"].value is None
    assert ws["B2"].number_format == "0.00"


def test_numbers_typed_as_text_are_stored_as_numbers(docs_dir):
    with _client() as c:
        c.post("/api/documents/create", json={"name": "calc", "kind": "sheet"})
        assert _put(c, "calc.xlsx", [[{"v": "3"}, {"v": "2.5"}, {"v": "nan"}, {"v": "x"}]]).status_code == 200
    ws = load_workbook(docs_dir / "calc.xlsx").active
    assert ws["A1"].value == 3 and isinstance(ws["A1"].value, int)
    assert ws["B1"].value == 2.5
    assert ws["C1"].value == "nan"             # not a float NaN
    assert ws["D1"].value == "x"


def test_a_save_to_a_new_name_still_makes_a_workbook(docs_dir):
    with _client() as c:
        assert _put(c, "fresh.xlsx", [[{"v": "hello"}]]).status_code == 200
    assert load_workbook(docs_dir / "fresh.xlsx").active["A1"].value == "hello"


def test_a_save_leaves_no_temporary_file_behind(docs_dir):
    _budget(docs_dir / "budget.xlsx")
    with _client() as c:
        assert _put(c, "budget.xlsx", _grid(c, "budget.xlsx")).status_code == 200
    assert sorted(p.name for p in docs_dir.iterdir()) == ["budget.xlsx"]


# ─── .xlsx parts openpyxl would drop: refuse, don't strip ───────────
def test_a_workbook_with_a_chart_is_refused_and_left_alone(docs_dir):
    path = docs_dir / "chart.xlsx"
    wb = Workbook()
    ws = wb.active
    for i in range(1, 4):
        ws.cell(row=i, column=1, value=i)
    chart = BarChart()
    chart.add_data(Reference(ws, min_col=1, min_row=1, max_row=3))
    ws.add_chart(chart, "C1")
    wb.save(path)
    before = path.read_bytes()
    with _client() as c:
        rows = _grid(c, "chart.xlsx")
        rows[0][0] = {"v": "9"}
        r = _put(c, "chart.xlsx", rows)
    assert r.status_code == 409
    assert "charts, pictures or pivot tables" in r.json()["detail"]
    assert path.read_bytes() == before


def test_cell_comments_alone_do_not_trigger_the_refusal(docs_dir):
    # Comments live in a vmlDrawing part, which openpyxl does keep.
    from openpyxl.comments import Comment

    path = docs_dir / "noted.xlsx"
    wb = Workbook()
    wb.active["A1"] = "x"
    wb.active["A1"].comment = Comment("a note", "me")
    wb.save(path)
    with _client() as c:
        rows = _grid(c, "noted.xlsx")
        rows[0][0] = {"v": "y"}
        assert _put(c, "noted.xlsx", rows).status_code == 200
    ws = load_workbook(path).active
    assert ws["A1"].value == "y"
    assert ws["A1"].comment is not None and ws["A1"].comment.text == "a note"


# ─── The read sends ragged rows, as .csv always did ─────────────────
def test_the_read_does_not_pad_rows_to_the_column_cap(docs_dir):
    wb = Workbook()
    wb.active.append(["a", "b"])
    wb.active.append(["c"])
    wb.save(docs_dir / "small.xlsx")
    with _client() as c:
        rows = _grid(c, "small.xlsx")
    assert [[cell["v"] for cell in row] for row in rows] == [["a", "b"], ["c"]]


# ─── .csv: what the editor never saw is carried over ────────────────
def test_csv_rows_and_columns_past_the_window_survive(docs_dir):
    wide = ",".join(f"c{i}" for i in range(70))
    lines = [wide] + [f"r{i}" for i in range(1, 1005)]
    (docs_dir / "big.csv").write_text("\n".join(lines) + "\n", encoding="utf-8")
    with _client() as c:
        rows = _grid(c, "big.csv")
        assert len(rows) == docs._SHEET_MAX_ROWS and len(rows[0]) == docs._SHEET_MAX_COLS
        rows[0][0] = {"v": "first"}
        assert _put(c, "big.csv", rows).status_code == 200
    out = (docs_dir / "big.csv").read_text(encoding="utf-8").splitlines()
    head = out[0].split(",")
    assert head[0] == "first" and head[69] == "c69" and len(head) == 70
    assert len(out) == 1005 and out[-1] == "r1004"


def test_csv_keeps_its_byte_order_mark(docs_dir):
    (docs_dir / "bom.csv").write_bytes(codecs.BOM_UTF8 + "a,b\n".encode("utf-8"))
    with _client() as c:
        assert _put(c, "bom.csv", [[{"v": "x"}, {"v": "y"}]]).status_code == 200
    raw = (docs_dir / "bom.csv").read_bytes()
    assert raw.startswith(codecs.BOM_UTF8) and raw[3:] == b"x,y\r\n"


def test_more_rows_than_the_window_is_refused(docs_dir):
    # The request model caps rows at the window (SheetWriteRequest), so the
    # carry-over above never has to reconcile a grid longer than the window.
    with _client() as c:
        r = _put(c, "huge.csv", [[{"v": "x"}]] * (docs._SHEET_MAX_ROWS + 1))
    assert r.status_code == 422
    assert not (docs_dir / "huge.csv").exists()
