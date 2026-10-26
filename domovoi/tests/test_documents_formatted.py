"""Formatted documents (.html): the allowlist, the file, the editor's routes.

The Documents editor's new format is one self-contained HTML page per
document (domovoi/richdoc). These files can come from anywhere, so:

* every load and every save goes through ``richdoc.sanitize_html`` ON THE
  SERVER: scripts, event handlers, ``javascript:`` links, outside
  pictures and unknown style properties never survive, whatever the
  client sent;
* the file is a whole page that reads as a document in any browser, and
  the editor gets back only its content;
* .md opens and saves as markdown, .docx converts both ways, and
  nothing a save-as writes lands on top of an existing file.
"""

from __future__ import annotations

import base64
import io

import pytest
from fastapi.testclient import TestClient

from domovoi import richdoc
from domovoi.config import settings
from domovoi.richdoc import sanitize_html
from domovoi.richdoc.fileformat import BODY_END, BODY_START
from domovoi.richdoc.word import html_to_docx
from domovoi.tests.auth_testkit import install_fake_db
from web.backend.api import files as files_api
from web.backend.main import app

PNG_1PX = (
    "data:image/png;base64,iVBORw0KGgoAAAANSUhEUgAAAAEAAAABCAYAAAAfFcSJAAAADUlEQVR42mP8z8Dw"
    "HwAFBQIAX8jx0gAAAABJRU5ErkJggg=="
)


@pytest.fixture(autouse=True)
def _pre_setup(monkeypatch):
    install_fake_db(monkeypatch, admin=False)

    async def _not_blocked(device_id):
        return None

    monkeypatch.setattr(files_api, "_block_status", _not_blocked)


@pytest.fixture
def docs_dir(monkeypatch, tmp_path):
    monkeypatch.setattr(settings, "documents_dir", str(tmp_path))
    return tmp_path


def _client():
    return TestClient(app, headers={"X-Requested-With": "domovoi-tests"})


# ═══ The allowlist ═══════════════════════════════════════════════════

@pytest.mark.parametrize("payload", [
    "<script>alert(1)</script>",
    "<SCRIPT SRC=//evil/x.js></SCRIPT>",
    "<style>body{background:url(//evil)}</style>",
    "<iframe src='//evil'></iframe>",
    "<svg><script>alert(1)</script></svg>",
    "<svg onload=alert(1)></svg>",
    "<math><mi xlink:href='javascript:alert(1)'>x</mi></math>",
    "<object data='//evil'></object><embed src='//evil'>",
    "<template><img src=x onerror=alert(1)></template>",
    "<noscript><p title='</noscript><img src=x onerror=alert(1)>'></noscript>",
])
def test_active_content_goes_with_its_content(payload):
    out = sanitize_html(f"<p>before</p>{payload}<p>after</p>")
    assert out.startswith("<p>before</p>")
    assert out.endswith("<p>after</p>")
    low = out.lower()
    for word in ("script", "alert", "iframe", "onerror", "onload", "evil", "<svg", "<math", "<object"):
        assert word not in low, out


def test_an_unclosed_raw_text_element_takes_the_rest_not_the_page():
    # Fail closed: what follows an unclosed <svg>/<script> is never written.
    assert sanitize_html("<p>a</p><svg onload=alert(1)><p>b</p>") == "<p>a</p>"
    assert sanitize_html("<p>a</p><script>alert(1)") == "<p>a</p>"
    # ...but a void one (no end tag in HTML) takes only itself.
    assert sanitize_html("<p>a</p><embed src=//evil><link rel=x href=//evil><p>b</p>") == "<p>a</p><p>b</p>"


@pytest.mark.parametrize("href", [
    "javascript:alert(1)",
    "JaVaScRiPt:alert(1)",
    "java\nscript:alert(1)",
    "  javascript:alert(1)",
    "java&#115;cript:alert(1)",
    "&#106;avascript:alert(1)",
    "vbscript:msgbox(1)",
    "data:text/html,<script>alert(1)</script>",
])
def test_links_that_run_code_lose_their_address(href):
    out = sanitize_html(f'<p><a href="{href}">x</a></p>')
    assert "href" not in out and "script" not in out.lower().replace("noscript", "")
    assert ">x</a>" in out  # the words stay


def test_safe_links_stay_and_never_open_a_window_back():
    out = sanitize_html('<p><a href="https://example.org/a?b=1&c=2" target="_blank" rel="opener">x</a>'
                        '<a href="mailto:alex@example.org">m</a><a href="#top">t</a></p>')
    assert 'href="https://example.org/a?b=1&amp;c=2"' in out
    assert 'href="mailto:alex@example.org"' in out and 'href="#top"' in out
    assert out.count('rel="noopener noreferrer nofollow"') == 3
    assert "opener\"" not in out.replace("noopener", "")


def test_event_handlers_and_unknown_attributes_are_dropped():
    out = sanitize_html('<p onclick="x()" onmouseover=y() id="a" class="evil" data-x="1" '
                        'style="text-align:center">t</p>')
    assert out == '<p style="text-align: center">t</p>'


def test_only_embedded_raster_pictures_survive():
    assert sanitize_html(f'<img src="{PNG_1PX}" alt="dot">') == f'<img src="{PNG_1PX}" alt="dot">'
    # an outside picture would make opening the document fetch from the web
    assert sanitize_html('<img src="https://tracker.example/p.gif">') == ""
    assert sanitize_html('<img src="/api/documents/raw/x.png">') == ""
    # SVG can carry script; never as a picture here
    svg = "data:image/svg+xml;base64," + base64.b64encode(b"<svg onload='alert(1)'/>").decode()
    assert sanitize_html(f'<img src="{svg}">') == ""
    assert sanitize_html('<img src="data:text/html;base64,PHNjcmlwdD4=">') == ""


@pytest.mark.parametrize(("style", "kept"), [
    ("color: #ff0000; font-family: Inter, sans-serif; font-size: 18px", "color: #ff0000; font-family: Inter, sans-serif; font-size: 18px"),
    ("color: rgb(10, 20, 30); background-color: yellow", "color: rgb(10, 20, 30); background-color: yellow"),
    ("font-size: 12pt; position: fixed; top: 0", "font-size: 12pt"),
    ("color: red; background-image: url(https://evil/x)", None),
    ("background: url(javascript:alert(1))", None),
    ("width: expression(alert(1))", None),
    ("color: \\72 ed", None),
    ("color: red /* */", None),
    ("font-family: x; color: var(--a)", None),
    ("font-size: 99999px", None),
])
def test_style_keeps_only_what_the_editor_writes(style, kept):
    out = sanitize_html(f'<p><span style="{style}">t</span></p>')
    if kept is None:
        assert out == "<p><span>t</span></p>"
    else:
        assert out == f'<p><span style="{kept}">t</span></p>'


def test_style_properties_are_per_element():
    # a paragraph may be aligned, not coloured or positioned
    assert sanitize_html('<p style="color:red;text-align:right">t</p>') == '<p style="text-align: right">t</p>'
    assert sanitize_html('<figure style="width: 40%; float: left">x</figure>') == '<figure style="width: 40%">x</figure>'


def test_the_editors_own_structures_survive_unchanged():
    doc = (
        '<h2 style="text-align: center">Plan</h2>'
        '<ul data-type="taskList"><li data-type="taskItem" data-checked="true"><label>'
        '<input type="checkbox" checked><span></span></label><div><p>milk</p></div></li></ul>'
        '<table style="min-width: 75px"><colgroup><col style="width: 120px"></colgroup><tbody>'
        '<tr><th colspan="1" rowspan="1" colwidth="120"><p>a</p></th></tr></tbody></table>'
        f'<figure data-align="right" data-width="40" style="width: 40%"><img src="{PNG_1PX}" alt="x">'
        '<figcaption>cap</figcaption></figure>'
        '<p><mark data-color="#fff3a3" style="background-color: #fff3a3">hi</mark> <s>no</s> <u>u</u> '
        '<code>c</code></p><pre><code class="language-python">x = 1</code></pre><blockquote><p>q</p></blockquote><hr>'
    )
    assert sanitize_html(doc) == doc


def test_old_spellings_become_the_editors_own():
    assert sanitize_html("<p><b>b</b><i>i</i><strike>s</strike><del>d</del></p>") == \
        "<p><strong>b</strong><em>i</em><s>s</s><s>d</s></p>"


def test_output_is_balanced_and_text_is_escaped():
    out = sanitize_html("<p>a <strong>b <em>c</p><p>1 &lt; 2 &amp; <b>x")
    assert out == "<p>a <strong>b <em>c</em></strong></p><p>1 &lt; 2 &amp; <strong>x</strong></p>"
    assert sanitize_html("<p>a<p>b") == "<p>a</p><p>b</p>"


def test_sanitising_twice_changes_nothing():
    messy = ('<div><p style="color:red;text-align:left" onclick=x>a<b>b</div><ul><li>1<li>2</ul>'
             '<a href="javascript:x">l</a><table><tr><td>c<td>d</table>')
    once = sanitize_html(messy)
    assert sanitize_html(once) == once


def test_a_hostile_nesting_depth_is_flattened_not_followed():
    out = sanitize_html("<span>" * 5000 + "deep" + "</span>" * 5000)
    assert "deep" in out
    assert out.count("<span>") <= 120


def test_comments_doctypes_and_processing_instructions_are_dropped():
    out = sanitize_html("<!DOCTYPE html><!-- <script>x</script> --><?xml x?><p>t</p><![CDATA[ y ]]>")
    assert out == "<p>t</p>"


# ═══ The file ════════════════════════════════════════════════════════

def test_the_file_is_a_whole_page_and_gives_back_only_its_content():
    body = "<h1>Hi</h1><p>there</p>"
    page = richdoc.document_file(body, "Groceries & more")
    assert page.startswith("<!DOCTYPE html>")
    assert "<title>Groceries &amp; more</title>" in page
    assert "<style>" in page and "<link" not in page and "<script" not in page
    assert "@media print" in page  # prints like a document too
    assert richdoc.document_body(page) == body
    assert richdoc.looks_like_document(page)


def test_a_page_from_anywhere_else_is_read_from_its_body():
    foreign = ("<html><head><title>x</title><style>p{}</style><script>evil()</script></head>"
               "<body onload='evil()'><p>keep</p><script>evil()</script></body></html>")
    assert sanitize_html(richdoc.document_body(foreign)) == "<p>keep</p>"
    assert sanitize_html(richdoc.document_body("<p>bare</p>")) == "<p>bare</p>"


# ═══ The readers' text ═══════════════════════════════════════════════

def test_readers_get_markdown_shaped_text_without_pictures_bytes():
    md = richdoc.html_to_markdown(richdoc.document_file(
        '<h1>Lease</h1><p>Rent is <strong>due</strong>\non the 1st.</p>'
        '<ul data-type="taskList"><li data-type="taskItem" data-checked="false"><label><input type="checkbox"></label>'
        '<div><p>sign</p></div></li></ul><ol><li><p>one</p></li><li><p>two</p></li></ol>'
        '<table><tbody><tr><th><p>A</p></th><th><p>B</p></th></tr><tr><td><p>1</p></td><td><p>2</p></td></tr></tbody></table>'
        f'<figure><img src="{PNG_1PX}" alt="floor plan"><figcaption>The flat</figcaption></figure>'
        '<pre><code>a\n  b</code></pre><script>secret()</script>', "Lease"))
    assert md.split("\n\n") == [
        "# Lease", "Rent is due on the 1st.", "- [ ] sign", "1. one\n2. two",
        "| A | B |\n|---|---|\n| 1 | 2 |", "[Picture: floor plan]", "The flat", "```\na\n  b\n```\n",
    ]
    assert "base64" not in md and "secret" not in md


# ═══ Word ════════════════════════════════════════════════════════════

def test_word_keeps_the_structure_and_the_formatting():
    from docx import Document

    data = html_to_docx(sanitize_html(
        '<h1 style="text-align: center">Title</h1>'
        '<p><strong>bold</strong> <em>it</em> <span style="color: #ff0000; font-size: 18px; '
        'font-family: Georgia, serif">red</span> <a href="https://example.org">link</a></p>'
        '<ul><li><p>a</p><ul><li><p>a1</p></li></ul></li></ul><ol><li><p>n</p></li></ol>'
        '<ul data-type="taskList"><li data-type="taskItem" data-checked="true"><label><input type="checkbox" checked>'
        '</label><div><p>done</p></div></li></ul>'
        '<table><tbody><tr><th colspan="2"><p>H</p></th></tr><tr><td><p>x</p></td><td><p>y</p></td></tr></tbody></table>'
        f'<figure data-align="center" data-width="50"><img src="{PNG_1PX}" alt="dot"><figcaption>Cap</figcaption></figure>'
    ), "Title")
    d = Document(io.BytesIO(data))
    paras = [(p.style.name, p.text) for p in d.paragraphs]
    assert ("Heading 1", "Title") in paras
    assert ("List Bullet", "a") in paras and ("List Bullet 2", "a1") in paras
    assert ("List Number", "n") in paras
    assert any(t == "☑ done" for _, t in paras)
    assert ("Caption", "Cap") in paras
    runs = {r.text: r for p in d.paragraphs for r in p.runs}
    assert runs["bold"].bold and runs["it"].italic
    assert str(runs["red"].font.color.rgb) == "FF0000"
    assert runs["red"].font.size.pt == 13.5 and runs["red"].font.name == "Georgia"
    assert "example.org" in str([r.target_ref for r in d.part.rels.values()])
    table = d.tables[0]
    assert table.cell(0, 0).text == "H" and table.cell(0, 1).text == "H"  # merged
    assert table.cell(1, 1).text == "y"
    assert len(d.inline_shapes) == 1


# ═══ The routes ══════════════════════════════════════════════════════

def test_a_save_is_sanitised_on_the_server_whatever_the_client_sent(docs_dir):
    with _client() as c:
        r = c.put("/api/documents/doc/plan.html", json={
            "html": '<p onclick="x()">hi<script>steal()</script><img src="https://t.example/p.gif">'
                    '<a href="javascript:steal()">l</a></p>'})
    assert r.status_code == 200, r.text
    assert r.json()["rel_path"] == "plan.html" and r.json()["category"] == "doc"
    page = (docs_dir / "plan.html").read_text(encoding="utf-8")
    assert page.startswith("<!DOCTYPE html>") and "<title>plan</title>" in page
    body = page[page.index(BODY_START) + len(BODY_START): page.index(BODY_END)].strip()
    assert body == '<p>hi<a rel="noopener noreferrer nofollow">l</a></p>'


def test_a_load_is_sanitised_too_for_a_file_from_anywhere(docs_dir):
    (docs_dir / "downloaded.html").write_text(
        "<html><head><script>evil()</script></head><body><h1 onmouseover=evil()>Hi</h1>"
        "<iframe src=//evil></iframe><p>text</p></body></html>", encoding="utf-8")
    with _client() as c:
        r = c.get("/api/documents/doc/downloaded.html")
    assert r.status_code == 200, r.text
    j = r.json()
    assert j["format"] == "html" and j["name"] == "downloaded"
    assert j["html"] == "<h1>Hi</h1><p>text</p>"


def test_markdown_opens_and_saves_as_markdown(docs_dir):
    (docs_dir / "notes.md").write_bytes(b"# Hi\n\n- a\n")
    with _client() as c:
        got = c.get("/api/documents/doc/notes.md").json()
        assert got["format"] == "md" and got["markdown"] == "# Hi\n\n- a\n"
        assert c.put("/api/documents/doc/notes.md", json={"html": "<p>x</p>"}).status_code == 422
        r = c.put("/api/documents/doc/notes.md", json={"markdown": "# Hi\n\n- a\n- b\n"})
    assert r.status_code == 200
    assert (docs_dir / "notes.md").read_text(encoding="utf-8") == "# Hi\n\n- a\n- b\n"


def test_word_opens_by_details_and_saves_as_word(docs_dir):
    from docx import Document

    with _client() as c:
        c.post("/api/documents/create", json={"name": "letter.docx", "kind": "doc"})
        got = c.get("/api/documents/doc/letter.docx").json()
        assert got["format"] == "docx" and "html" not in got
        r = c.put("/api/documents/doc/letter.docx", json={"html": "<h1>Dear Alex</h1><p>Hello</p>"})
    assert r.status_code == 200, r.text
    assert [p.text for p in Document(docs_dir / "letter.docx").paragraphs] == ["Dear Alex", "Hello"]


def test_a_save_never_overwrites_what_the_editor_could_not_have_opened(docs_dir):
    (docs_dir / "fake.docx").write_bytes(b"not a zip at all")
    (docs_dir / "bin.md").write_bytes(b"\xff\xfe\x00binary")
    with _client() as c:
        assert c.put("/api/documents/doc/fake.docx", json={"html": "<p>x</p>"}).status_code == 409
        assert c.put("/api/documents/doc/bin.md", json={"markdown": "x"}).status_code == 415
        for name in ("tax.pdf", "budget.xlsx", "notes.txt", "page.svg"):
            assert c.put(f"/api/documents/doc/{name}", json={"html": "<p>x</p>"}).status_code == 415
    assert (docs_dir / "fake.docx").read_bytes() == b"not a zip at all"
    assert not (docs_dir / "tax.pdf").exists()


def test_the_text_editor_cannot_write_a_formatted_document_raw(docs_dir):
    """PUT /text would skip the allowlist: a formatted document is written
    only through PUT /doc (and a .docx is not text at all)."""
    (docs_dir / "plan.html").write_text(richdoc.document_file("<p>a</p>", "plan"), encoding="utf-8")
    before = (docs_dir / "plan.html").read_bytes()
    with _client() as c:
        for name in ("plan.html", "new.htm", "letter.docx"):
            r = c.put(f"/api/documents/text/{name}", json={"text": "<script>x()</script>"})
            assert r.status_code == 415, (name, r.text)
    assert (docs_dir / "plan.html").read_bytes() == before
    assert not (docs_dir / "new.htm").exists() and not (docs_dir / "letter.docx").exists()


def test_save_as_makes_a_new_file_beside_the_old_one_and_never_on_top(docs_dir):
    (docs_dir / "notes.md").write_text("# Hi\n", encoding="utf-8")
    (docs_dir / "notes.html").write_text("someone else's", encoding="utf-8")
    with _client() as c:
        r = c.post("/api/documents/doc-save-as", json={
            "from_path": "notes.md", "format": "html", "html": "<h1>Hi</h1>"})
        assert r.status_code == 200, r.text
        assert r.json()["rel_path"] == "notes (1).html"
        w = c.post("/api/documents/doc-save-as", json={
            "from_path": "notes.md", "format": "docx", "html": "<h1>Hi</h1>"})
        assert w.json()["rel_path"] == "notes.docx"
    assert (docs_dir / "notes.html").read_text(encoding="utf-8") == "someone else's"
    assert (docs_dir / "notes.md").read_text(encoding="utf-8") == "# Hi\n"
    assert "<h1>Hi</h1>" in (docs_dir / "notes (1).html").read_text(encoding="utf-8")


def test_save_as_stays_in_the_documents_folder(docs_dir):
    with _client() as c:
        r = c.post("/api/documents/doc-save-as", json={
            "from_path": "../outside.md", "format": "html", "html": "<p>x</p>"})
    assert r.status_code == 400
    assert not (docs_dir.parent / "outside.html").exists()


def test_download_as_word_converts_what_is_on_screen(docs_dir):
    with _client() as c:
        r = c.post("/api/documents/convert/docx", json={"html": "<p>Grüße<script>x</script></p>",
                                                        "title": "Brief für Oma"})
    assert r.status_code == 200
    assert "wordprocessingml" in r.headers["content-type"]
    assert "filename*=UTF-8''Brief%20f%C3%BCr%20Oma.docx" in r.headers["content-disposition"]
    from docx import Document

    assert [p.text for p in Document(io.BytesIO(r.content)).paragraphs] == ["Grüße"]


def test_a_formatted_document_exports_to_word_from_disk(docs_dir):
    (docs_dir / "p.html").write_text(richdoc.document_file("<h2>Two</h2>", "p"), encoding="utf-8")
    with _client() as c:
        r = c.get("/api/documents/export/doc/p.html", params={"fmt": "docx"})
    assert r.status_code == 200
    from docx import Document

    assert [p.text for p in Document(io.BytesIO(r.content)).paragraphs] == ["Two"]


def test_the_raw_file_is_still_never_rendered_on_the_dashboards_origin(docs_dir):
    (docs_dir / "p.html").write_text(richdoc.document_file("<p>x</p>", "p"), encoding="utf-8")
    with _client() as c:
        r = c.get("/api/documents/raw/p.html")
    assert r.status_code == 200
    assert r.headers["content-disposition"].startswith("attachment")
    assert "sandbox" in r.headers.get("content-security-policy", "")


def test_the_editor_routes_have_room_for_a_document_with_pictures():
    from domovoi.transport_guard import body_limit_for, default_body_limit

    for path in ("/api/documents/doc/a.html", "/api/documents/doc-save-as", "/api/documents/convert/docx"):
        assert body_limit_for(path) >= 40 * 1024 * 1024 > default_body_limit()


# ═══ Everything that reads documents reads these too ═════════════════

LEASE = richdoc.document_file(
    "<h1>Lease</h1><p>The flat at <strong>12 Elm Road</strong>.</p><h2>Deposit</h2>"
    "<p>Two months' rent, returned within 30 days.</p>"
    f'<figure><img src="{PNG_1PX}" alt="floor plan"><figcaption>Floor plan</figcaption></figure>'
    "<script>steal()</script>", "Lease")


def test_chat_reads_a_formatted_document_by_its_sections(tmp_path):
    from domovoi import doc_extract

    path = tmp_path / "Lease.html"
    path.write_text(LEASE, encoding="utf-8")
    assert doc_extract.kind_for("Lease.html") == "html" and doc_extract.kind_for("x.HTM") == "html"
    assert doc_extract.sniff(path.read_bytes()[:8192], "html") is None
    got = doc_extract.extract(path, "html")
    text = " ".join(f"{p.get('head') or ''} {p.get('text') or ''}" for p in got["parts"])
    labels = [p.get("label") for p in got["parts"]]
    assert "Section: Deposit" in labels
    assert "12 Elm Road" in text and "returned within 30 days" in text and "Floor plan" in text
    assert "base64" not in text and "steal" not in text and "<" not in text


def test_the_existing_markdown_reading_is_unchanged(tmp_path):
    from domovoi import doc_extract

    path = tmp_path / "notes.md"
    path.write_bytes(b"# Notes\n\nMilk and eggs.\n")
    got = doc_extract.extract(path, "md")
    assert got["kind"] == "md" and "Milk and eggs." in got["parts"][0]["text"]


def test_read_aloud_takes_formatted_documents():
    from domovoi.voice_docs import reading

    assert {".html", ".htm", ".md", ".pdf"} <= reading.READABLE_EXTS


def test_a_formatted_document_prints_as_its_text(tmp_path):
    from domovoi.printing import jobs, render_worker

    assert jobs.kind_of("Lease.html") == "document"
    assert ".html" in jobs.printable_info()["exts"]
    (tmp_path / "input").write_text(LEASE, encoding="utf-8")
    out = render_worker.prepare({"kind": "document", "paper": [612, 792], "title": "Lease"}, tmp_path)
    pdf = (tmp_path / "doc.pdf").read_bytes()
    assert pdf.startswith(b"%PDF-")
    assert "without pictures" in " ".join(out.get("notes") or [])


def test_ask_about_this_document_is_offered_for_formatted_documents():
    from pathlib import Path

    jsx = (Path(__file__).resolve().parents[2] / "web" / "static" / "chat_documents.jsx").read_text(encoding="utf-8")
    reads = next(line for line in jsx.splitlines() if "const READS" in line)
    assert "'.html'" in reads and "'.md'" in reads


def test_print_and_pdf_use_the_same_page_style_as_the_saved_file():
    """"Download → PDF" prints the page the editor builds in the browser
    (doc_editor.jsx PRINT_CSS); a saved file carries fileformat.STYLE.
    They are one look, so they are one text."""
    import re
    from pathlib import Path

    from domovoi.richdoc.fileformat import STYLE

    jsx = (Path(__file__).resolve().parents[2] / "web" / "static" / "doc_editor.jsx").read_text(encoding="utf-8")
    m = re.search(r"const PRINT_CSS = `(.*?)`;", jsx, re.S)
    assert m, "PRINT_CSS not found in doc_editor.jsx"
    assert m.group(1).strip() == STYLE.strip()


def test_every_editor_save_tells_co_editing(docs_dir, monkeypatch):
    """wf/doc-collab: a save that doesn't call the hook looks like an outside
    change, and the shared document is rebuilt from the file."""
    from web.backend.api import documents as docs_api

    told = []

    async def _saved(request, target):
        told.append(target.name)

    monkeypatch.setattr(docs_api, "_collab_saved", _saved)
    (docs_dir / "n.md").write_bytes(b"# n\n")
    with _client() as c:
        assert c.put("/api/documents/doc/a.html", json={"html": "<p>x</p>"}).status_code == 200
        assert c.post("/api/documents/doc-save-as", json={
            "from_path": "n.md", "format": "html", "html": "<p>n</p>"}).status_code == 200
    assert told == ["a.html", "n.html"]
