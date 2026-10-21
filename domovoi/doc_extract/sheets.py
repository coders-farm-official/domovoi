"""Spreadsheets (.xlsx, .csv) as parts of 40 rows, the header row repeated.

* **.xlsx**: every sheet, through openpyxl ``read_only`` (MIT, already a
  dependency), twice: the saved values, and the formulas where a value is
  empty (a file saved without calculating, as the phone's sheet editor
  writes them, shows ``=SUM(B2:B9)``; D0: 54/70 cells, as expected). The
  zip is checked first (:class:`~domovoi.doc_extract.safe.SafeZip`) and
  every XML part in it goes through the entity-refusing parser before
  openpyxl opens it.
* **.csv**: the delimiter sniffed (comma, semicolon, tab), a UTF-8 BOM
  dropped, cp1252 when it isn't UTF-8 (D0: 263/263 cells).
* **Numbers from the app, not the model**: each sheet gets a "column
  totals" part, computed here for every numeric column (how many
  numbers, the total, the average, the lowest and the highest), so "what's
  the total?" is quoted, not calculated.
"""

from __future__ import annotations

import csv
import datetime as dt
import io
from decimal import Decimal, InvalidOperation
from pathlib import Path
from typing import Any

from domovoi.doc_extract.safe import Refused, SafeZip, check_xml
from domovoi.doc_extract.structure import md_table
from domovoi.doc_qa.text import clean as normalize

ROWS_PER_PART = 40
#: A column is numeric when this share of its filled cells are numbers.
NUMERIC_SHARE = 0.6


def cell_text(v: Any) -> str:
    if v is None:
        return ""
    if isinstance(v, dt.datetime):
        return v.strftime("%Y-%m-%d %H:%M") if (v.hour or v.minute) else v.strftime("%Y-%m-%d")
    if isinstance(v, (dt.date, dt.time)):
        return v.isoformat()
    if isinstance(v, bool):
        return "TRUE" if v else "FALSE"
    if isinstance(v, float):
        if v != v or v in (float("inf"), float("-inf")):
            return str(v)
        return str(int(v)) if v == int(v) and abs(v) < 1e15 else f"{v:.2f}".rstrip("0").rstrip(".")
    return normalize(str(v))


def number_of(s: str) -> Decimal | None:
    """The number a cell shows ("$1,200.50", "(40)", "12%"), or None."""
    t = s.strip().replace(" ", "").replace(" ", "")
    if not t:
        return None
    neg = t.startswith("(") and t.endswith(")")
    t = t.strip("()")
    for sym in ("$", "€", "£", "¥", "%"):
        t = t.replace(sym, "")
    t = t.replace(",", "")
    if t.startswith("-"):
        neg, t = (not neg), t[1:]
    if not t or not (t.replace(".", "", 1).isdigit()):
        return None
    try:
        d = Decimal(t)
    except InvalidOperation:
        return None
    return -d if neg else d


def _fmt(d: Decimal) -> str:
    q = d.quantize(Decimal("0.01")) if d != d.to_integral_value() else d.quantize(Decimal("1"))
    s = f"{q:,}"
    return s


def column_totals(header: list[str], rows: list[list[str]]) -> list[str]:
    lines = []
    width = max([len(header)] + [len(r) for r in rows]) if rows else len(header)
    for c in range(width):
        filled = [r[c] for r in rows if c < len(r) and r[c].strip()]
        nums = [n for n in (number_of(x) for x in filled) if n is not None]
        if len(nums) < 2 or len(nums) < NUMERIC_SHARE * len(filled):
            continue
        name = (header[c].strip() if c < len(header) and header[c].strip() else f"column {c + 1}")
        total = sum(nums, Decimal(0))
        avg = total / len(nums)
        lines.append(f"{name}: total {_fmt(total)}, average {_fmt(avg)}, lowest {_fmt(min(nums))}, "
                     f"highest {_fmt(max(nums))} ({len(nums)} numbers)")
    return lines


def _md_row(r: list[str], width: int) -> str:
    cells = [" ".join(c.replace("|", "/").split()) for c in r] + [""] * (width - len(r))
    return "| " + " | ".join(cells) + " |"


def sheet_parts(rows: list[list[str]], sheet: str | None) -> list[dict[str, Any]]:
    """One sheet's rows (row 1 first) as parts of :data:`ROWS_PER_PART` rows,
    labelled by their row numbers, plus its column totals."""
    numbered = [(i + 1, r) for i, r in enumerate(rows)]
    while numbered and not any(c.strip() for c in numbered[-1][1]):
        numbered.pop()
    numbered = [(n, r) for n, r in numbered]
    first = next((k for k, (_n, r) in enumerate(numbered) if any(c.strip() for c in r)), None)
    if first is None:
        return []
    header_no, header = numbered[first]
    data = [(n, r) for n, r in numbered[first + 1:] if any(c.strip() for c in r)]
    prefix = f"Sheet {sheet}, " if sheet else ""
    head = md_table([header]) if data else ""
    parts: list[dict[str, Any]] = []
    if not data:
        return [{"label": f"{prefix}row {header_no}".strip(), "text": md_table([header])}]
    width = max(len(header), *(len(r) for _n, r in data))
    for k in range(0, len(data), ROWS_PER_PART):
        chunk = data[k:k + ROWS_PER_PART]
        body = [_md_row(r, width) for _n, r in chunk]
        lo, hi = chunk[0][0], chunk[-1][0]
        label = f"{prefix}rows {lo}-{hi}" if hi != lo else f"{prefix}row {lo}"
        parts.append({"label": label, "head": head, "text": "\n".join(body)})
    totals = column_totals(header, [r for _n, r in data])
    if totals:
        where = f"sheet {sheet}" if sheet else "this file"
        parts.append({"label": f"{prefix}column totals",
                      "text": f"Column totals for {where}, worked out by Domovoi from rows "
                              f"{data[0][0]}-{data[-1][0]}:\n" + "\n".join(totals)})
    return parts


def extract_xlsx(path: Path) -> dict[str, Any]:
    with SafeZip(path, "spreadsheet") as z:
        if not z.has("xl/workbook.xml"):
            raise Refused("That isn't an Excel workbook Domovoi can read (no workbook part inside).")
        for name in z.xml_members():
            check_xml(z.read(name), "spreadsheet")
    from openpyxl import load_workbook

    try:
        vals = load_workbook(str(path), read_only=True, data_only=True)
        forms = load_workbook(str(path), read_only=True, data_only=False)
    except Exception as e:  # noqa: BLE001 — a damaged workbook
        raise Refused("That spreadsheet is damaged and couldn't be read.") from e
    parts: list[dict[str, Any]] = []
    no_value = 0
    sheets = []
    try:
        for ws in vals.worksheets:
            fws = forms[ws.title]
            rows: list[list[str]] = []
            for vr, fr in zip(ws.iter_rows(values_only=True), fws.iter_rows(values_only=True)):
                row = []
                for v, f in zip(vr, fr):
                    if v is None and isinstance(f, str) and f.startswith("="):
                        row.append(f)
                        no_value += 1
                    else:
                        row.append(cell_text(v))
                while row and row[-1] == "":
                    row.pop()
                rows.append(row)
            sheets.append(ws.title)
            parts += sheet_parts(rows, normalize(ws.title))
    finally:
        vals.close()
        forms.close()
    return {"parts": parts, "sheets": sheets, "formulas_without_value": no_value}


def decode_text(raw: bytes) -> tuple[str, str]:
    """(text, encoding): a UTF-8 BOM dropped, UTF-16 by its BOM, then
    UTF-8, then cp1252."""
    if raw.startswith(b"\xef\xbb\xbf"):
        return raw[3:].decode("utf-8", "replace"), "utf-8-sig"
    if raw.startswith((b"\xff\xfe", b"\xfe\xff")):
        return raw.decode("utf-16", "replace"), "utf-16"
    try:
        return raw.decode("utf-8"), "utf-8"
    except UnicodeDecodeError:
        return raw.decode("cp1252", "replace"), "cp1252"


def extract_csv(path: Path) -> dict[str, Any]:
    text, enc = decode_text(path.read_bytes())
    try:
        delim = csv.Sniffer().sniff(text[:8192], delimiters=",;\t").delimiter
    except csv.Error:
        delim = ","
    try:
        rows = [[normalize(c) for c in r] for r in csv.reader(io.StringIO(text), delimiter=delim)]
    except csv.Error as e:
        raise Refused("That CSV file is damaged and couldn't be read.") from e
    return {"parts": sheet_parts(rows, None), "encoding": enc, "delimiter": delim}
